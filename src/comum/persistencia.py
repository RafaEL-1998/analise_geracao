"""Gravação dos dados das etapas com cópia de segurança (spec do Tratamento de dados, FR-028 a FR-032; constituição, Requisito
Técnico 3).

Toda gravação de dados das etapas (``data/usinas/<slug>/<etapa>/``) passa por este módulo. A cópia
``<nome>.bak`` só fica nos dados da Coleta e do Tratamento; a Conferência e as Análises gravam com
``copia=False``, e a versão anterior serve só à restauração em falha:

1. o conteúdo novo é gravado em ``<nome>.tmp`` na mesma pasta e conferido (registros e colunas);
2. se for idêntico ao arquivo atual, nada é gravado e a cópia ``<nome>.bak`` fica como estava;
3. caso contrário, o arquivo atual é copiado para ``<nome>.bak.tmp``, o temporário substitui o
   destino (troca atômica), o SHA-256 do destino é conferido e só então a cópia vira ``<nome>.bak``;
4. em qualquer falha o destino volta à versão anterior, os temporários são removidos e
   ``ErroPersistencia`` é levantado.

Os formatos determinísticos (CSV, texto, Parquet) são comparados byte a byte. As planilhas mudam
de bytes a cada gravação (datas internas do pacote) e são comparadas pela assinatura dos dados,
gravada como propriedade personalizada ``assinatura_dados``.
"""

from __future__ import annotations

import csv
import hashlib
import re
import shutil
import zipfile
from contextlib import contextmanager
from enum import Enum
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, Iterator, List, Optional, Sequence, Tuple

import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook
from openpyxl.packaging.custom import StringProperty

from src.comum.logger import ONSError, setup_logger

logger = setup_logger("persistencia")

PROPRIEDADE_ASSINATURA = "assinatura_dados"
_RE_ASSINATURA = re.compile(
    rf'name="{PROPRIEDADE_ASSINATURA}"[^>]*>\s*<vt:lpwstr>([0-9a-f]{{64}})</vt:lpwstr>', re.S
)
_OPCOES_CSV_PADRAO: Dict[str, Any] = {
    "sep": ";",
    "index": False,
    "encoding": "utf-8",
    "date_format": "%Y-%m-%d %H:%M:%S",
}


class ResultadoGravacao(str, Enum):
    """Resultado de uma gravação de dados das etapas."""

    NOVO = "NOVO"
    ALTERADO = "ALTERADO"
    INALTERADO = "INALTERADO"


class ErroPersistencia(ONSError):
    """Falha ao gravar um arquivo processado; o arquivo foi mantido ou restaurado na versão anterior."""

    def __init__(self, destino: Path, motivo: str) -> None:
        super().__init__(f"Falha ao gravar {destino}: {motivo}")
        self.destino = Path(destino)
        self.motivo = motivo


# Registros ativos de gravações (``registrar_gravacoes``): cada etapa sabe o que gravou e com que resultado
_REGISTROS: List[List[Tuple[Path, ResultadoGravacao]]] = []


@contextmanager
def registrar_gravacoes() -> Iterator[List[Tuple[Path, ResultadoGravacao]]]:
    """Lista, na ordem, (arquivo, resultado) de cada gravação feita dentro do bloco."""
    registro: List[Tuple[Path, ResultadoGravacao]] = []
    _REGISTROS.append(registro)
    try:
        yield registro
    finally:
        _REGISTROS.remove(registro)


def _registrar(destino: Path, resultado: ResultadoGravacao) -> ResultadoGravacao:
    for registro in _REGISTROS:
        registro.append((destino, resultado))
    return resultado


# ---------------------------------------------------------------------------
# Núcleo
# ---------------------------------------------------------------------------


def _sha256(caminho: Path) -> str:
    h = hashlib.sha256()
    with open(caminho, "rb") as f:
        for bloco in iter(lambda: f.read(1024 * 1024), b""):
            h.update(bloco)
    return h.hexdigest()


def _bytes_iguais(atual: Path, novo: Path) -> bool:
    """Comparação byte a byte, sem cache (``filecmp.cmp`` reaproveita resultados com mesmo tamanho e data)."""
    if atual.stat().st_size != novo.stat().st_size:
        return False
    with open(atual, "rb") as a, open(novo, "rb") as b:
        while True:
            bloco_a, bloco_b = a.read(1024 * 1024), b.read(1024 * 1024)
            if bloco_a != bloco_b:
                return False
            if not bloco_a:
                return True


def _remover(*caminhos: Path) -> None:
    for caminho in caminhos:
        try:
            caminho.unlink(missing_ok=True)
        except OSError:
            logger.warning("Não foi possível remover o temporário %s", caminho)


def _gravar_com_copia(
    destino: Path,
    escrever: Callable[[Path], None],
    conferir: Callable[[Path], None],
    identico: Callable[[Path, Path], bool],
    descricao: str,
    copia_seguranca: bool = True,
) -> ResultadoGravacao:
    """Grava ``destino`` com conferência, comparação, cópia de segurança e restauração em falha.

    Com ``copia_seguranca`` falso, a versão anterior só serve à restauração em falha e não fica como ``.bak``.
    """
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    temporario = destino.with_name(destino.name + ".tmp")
    copia = destino.with_name(destino.name + ".bak")
    copia_candidata = destino.with_name(destino.name + ".bak.tmp")
    existia = destino.exists()
    trocado = False
    try:
        escrever(temporario)
        conferir(temporario)
        if existia and identico(destino, temporario):
            _remover(temporario)
            logger.info("INALTERADO: %s (%s)", destino.name, descricao)
            return _registrar(destino, ResultadoGravacao.INALTERADO)
        sha_novo = _sha256(temporario)
        if existia:
            shutil.copy2(destino, copia_candidata)
        temporario.replace(destino)
        trocado = True
        if _sha256(destino) != sha_novo:
            raise OSError("o conteúdo gravado no destino difere do conteúdo conferido")
        if existia:
            if copia_seguranca:
                copia_candidata.replace(copia)
            else:
                _remover(copia_candidata)
    except BaseException as exc:  # inclui a interrupção pelo usuário (Ctrl+C)
        if trocado:
            if existia and copia_candidata.exists():
                shutil.copy2(copia_candidata, destino)
            elif not existia:
                _remover(destino)
        _remover(temporario, copia_candidata)
        if not isinstance(exc, Exception):
            logger.error("Gravação de %s interrompida; o arquivo foi mantido na versão anterior.", destino)
            raise
        logger.error(
            "Falha ao gravar %s: %s. O arquivo foi mantido na versão anterior (verifique se ele está aberto em "
            "outro programa).", destino, exc,
        )
        raise ErroPersistencia(destino, str(exc)) from exc
    resultado = ResultadoGravacao.ALTERADO if existia else ResultadoGravacao.NOVO
    if existia and copia_seguranca:
        logger.info("ALTERADO: %s (%s); versão anterior em %s", destino.name, descricao, copia.name)
    elif existia:
        logger.info("ALTERADO: %s (%s)", destino.name, descricao)
    else:
        logger.info("NOVO: %s (%s)", destino.name, descricao)
    return _registrar(destino, resultado)


# ---------------------------------------------------------------------------
# CSV a partir de DataFrame
# ---------------------------------------------------------------------------


def _conferir_csv(caminho: Path, tabela: pd.DataFrame, opcoes: Dict[str, Any]) -> None:
    if len(tabela.columns) == 0:
        if not caminho.exists():
            raise OSError("arquivo temporário não foi criado")
        return
    lido = pd.read_csv(caminho, sep=opcoes.get("sep", ";"), encoding=opcoes.get("encoding", "utf-8"),
                       skip_blank_lines=False, dtype=str, keep_default_na=False)
    if lido.shape != tabela.shape:
        raise ValueError(f"releitura com {lido.shape} (linhas, colunas); esperado {tabela.shape}")


def gravar_csv(tabela: pd.DataFrame, destino: Path, copia: bool = True, **opcoes: Any) -> ResultadoGravacao:
    """Grava uma tabela em CSV (padrão: ';', UTF-8, sem índice, datas ISO); ``copia`` falso dispensa o ``.bak``."""
    opcoes = {**_OPCOES_CSV_PADRAO, **opcoes}
    return _gravar_com_copia(
        Path(destino),
        lambda tmp: tabela.to_csv(tmp, **opcoes),
        lambda tmp: _conferir_csv(tmp, tabela, opcoes),
        _bytes_iguais,
        f"{len(tabela)} linhas",
        copia_seguranca=copia,
    )


def gravar_bytes(conteudo: bytes, destino: Path, copia: bool = True) -> ResultadoGravacao:
    """Grava bytes (ex.: resultados serializados) com conferência; ``copia`` falso dispensa o ``.bak``."""

    def conferir(tmp: Path) -> None:
        if tmp.read_bytes() != conteudo:
            raise ValueError("bytes gravados diferem dos pretendidos")

    return _gravar_com_copia(Path(destino), lambda tmp: tmp.write_bytes(conteudo), conferir, _bytes_iguais,
                             f"{len(conteudo)} bytes", copia_seguranca=copia)


# ---------------------------------------------------------------------------
# CSV linha a linha (consolidação e auditoria da varredura)
# ---------------------------------------------------------------------------


def _conferir_linhas_csv(caminho: Path, cabecalho: Sequence[str], registros: int, delimitador: str) -> None:
    with open(caminho, "r", newline="", encoding="utf-8") as f:
        leitor = csv.reader(f, delimiter=delimitador)
        primeira = next(leitor, None)
        total = sum(1 for _ in leitor)
    if primeira != [str(c) for c in cabecalho]:
        raise ValueError("cabeçalho gravado difere do esperado")
    if total != registros:
        raise ValueError(f"{total} linhas gravadas; esperado {registros}")


def gravar_linhas_csv(
    cabecalho: Sequence[str],
    linhas: Iterable[Sequence[Any]],
    destino: Path,
    delimitador: str = ";",
) -> ResultadoGravacao:
    """Grava cabeçalho e linhas com o módulo csv (mesmo formato de bytes de ``csv.writer``)."""
    contagem = {"registros": 0}

    def escrever(tmp: Path) -> None:
        with open(tmp, "w", newline="", encoding="utf-8") as f:
            escritor = csv.writer(f, delimiter=delimitador)
            escritor.writerow(cabecalho)
            for linha in linhas:
                escritor.writerow(linha)
                contagem["registros"] += 1

    return _gravar_com_copia(
        Path(destino),
        escrever,
        lambda tmp: _conferir_linhas_csv(tmp, cabecalho, contagem["registros"], delimitador),
        _bytes_iguais,
        "linhas gravadas com o módulo csv",
    )


# ---------------------------------------------------------------------------
# Parquet
# ---------------------------------------------------------------------------


def _conferir_parquet(caminho: Path, tabela: pd.DataFrame) -> None:
    meta = pq.read_metadata(caminho)
    if (meta.num_rows, meta.num_columns) != tabela.shape:
        raise ValueError(f"Parquet com {(meta.num_rows, meta.num_columns)}; esperado {tabela.shape}")


def gravar_parquet(tabela: pd.DataFrame, destino: Path) -> ResultadoGravacao:
    """Grava uma tabela em Parquet (pyarrow, sem índice); a gravação é determinística."""
    return _gravar_com_copia(
        Path(destino),
        lambda tmp: tabela.to_parquet(tmp, engine="pyarrow", index=False),
        lambda tmp: _conferir_parquet(tmp, tabela),
        _bytes_iguais,
        f"{len(tabela)} linhas",
    )


# ---------------------------------------------------------------------------
# Planilha xlsx
# ---------------------------------------------------------------------------


def assinatura_dados(abas: Dict[str, pd.DataFrame]) -> str:
    """SHA-256 de Σ (nome da aba + "\\n" + CSV da aba sem índice), nas abas em ordem."""
    h = hashlib.sha256()
    for nome, tabela in abas.items():
        h.update(f"{nome}\n".encode("utf-8"))
        h.update(tabela.to_csv(index=False, sep=";", lineterminator="\n",
                               date_format="%Y-%m-%d %H:%M:%S").encode("utf-8"))
    return h.hexdigest()


def ler_assinatura(caminho: Path) -> Optional[str]:
    """Assinatura gravada na planilha (lida de docProps/custom.xml, sem abrir as abas)."""
    try:
        with zipfile.ZipFile(caminho) as pacote:
            xml = pacote.read("docProps/custom.xml").decode("utf-8")
    except (KeyError, OSError, zipfile.BadZipFile):
        return None
    achado = _RE_ASSINATURA.search(xml)
    return achado.group(1) if achado else None


def _conferir_planilha(caminho: Path, abas: Dict[str, pd.DataFrame], assinatura: str) -> None:
    with open(caminho, "rb") as f:
        livro = load_workbook(f, read_only=True)
        try:
            if livro.sheetnames != list(abas):
                raise ValueError(f"abas gravadas {livro.sheetnames}; esperado {list(abas)}")
            for nome, tabela in abas.items():
                folha = livro[nome]
                if len(tabela.columns) and (folha.max_row, folha.max_column) != (len(tabela) + 1, len(tabela.columns)):
                    raise ValueError(f"aba {nome} com {(folha.max_row, folha.max_column)} células; "
                                     f"esperado {(len(tabela) + 1, len(tabela.columns))}")
        finally:
            livro.close()
    if ler_assinatura(caminho) != assinatura:
        raise ValueError("assinatura dos dados ausente na planilha gravada")


def gravar_planilha(
    abas: Dict[str, pd.DataFrame],
    destino: Path,
    formatar: Optional[Callable[[pd.ExcelWriter], None]] = None,
) -> ResultadoGravacao:
    """Grava uma planilha com as abas na ordem dada; ``formatar(writer)`` ajusta larguras e painéis."""
    assinatura = assinatura_dados(abas)

    def escrever(tmp: Path) -> None:
        with open(tmp, "wb") as f:
            with pd.ExcelWriter(f, engine="openpyxl") as writer:
                for nome, tabela in abas.items():
                    tabela.to_excel(writer, sheet_name=nome, index=False)
                if formatar is not None:
                    formatar(writer)
                writer.book.custom_doc_props.append(StringProperty(name=PROPRIEDADE_ASSINATURA, value=assinatura))

    return _gravar_com_copia(
        Path(destino),
        escrever,
        lambda tmp: _conferir_planilha(tmp, abas, assinatura),
        lambda atual, _novo: ler_assinatura(atual) == assinatura,
        f"{len(abas)} abas, {sum(len(t) for t in abas.values())} linhas",
    )


# ---------------------------------------------------------------------------
# Texto (Markdown)
# ---------------------------------------------------------------------------


def _conferir_texto(caminho: Path, texto: str, newline: Optional[str]) -> None:
    with open(caminho, "r", encoding="utf-8", newline=None if newline is None else "") as f:
        lido = f.read()
    esperado = texto if newline is None else texto.replace("\n", newline)
    if lido != esperado:
        raise ValueError("texto relido difere do texto gravado")


def gravar_texto(texto: str, destino: Path, newline: Optional[str] = None) -> ResultadoGravacao:
    """Grava texto UTF-8 (fim de linha do sistema, como ``open(..., "w")``, salvo ``newline`` explícito)."""

    def escrever(tmp: Path) -> None:
        with open(tmp, "w", encoding="utf-8", newline=newline) as f:
            f.write(texto)

    return _gravar_com_copia(
        Path(destino),
        escrever,
        lambda tmp: _conferir_texto(tmp, texto, newline),
        _bytes_iguais,
        f"{texto.count(chr(10))} linhas de texto",
    )
