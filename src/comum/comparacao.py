"""Comparação do relatório com uma versão de referência (spec da Geração do relatório, ferramenta ``comparar``).

PDF, Markdown, CSV e figuras são comparados byte a byte; a planilha, célula a célula, aba a aba e na mesma ordem de
abas, porque os bytes do XLSX mudam a cada gravação sem mudar o conteúdo. Células numéricas são iguais quando
coincidem até a 12ª casa relativa: a diferença abaixo disso é resto de ponto flutuante (por exemplo, um número lido
de um CSV pelo leitor padrão do pandas), não muda nenhum texto nem nenhuma figura. O ``etapa.json`` fica fora da
comparação. A comparação não grava nada.

Execução: ``python -m src comparar --usina <slug> --referencia <pasta>``, com código 0 sem diferença, 6 com diferença
e 1 em erro.

``python -m src comparar --todas [--coleta]`` (spec 006, decisão R22; contrato cli-referencias.md): para cada relatório
de ``relatorios_referencia/``, refaz o Tratamento, a Conferência, as Análises e o Relatório num espaço isolado, a partir
da Coleta e do perfil congelados e com a data de geração da referência, e compara com o relatório aprovado. Não toca em
``data/usinas/``, ``reports/`` nem ``data/raw/``. Com ``--coleta``, refaz também a Coleta com ``--sem-portal`` e o
perfil congelado, noutro espaço isolado, e a compara com a congelada: as séries só no período da referência, as
auditorias só nas contagens por arquivo e os registros informativos sem o código 6.
"""

from __future__ import annotations

import math
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src import pipeline
from src.comum import caminhos, referencias
from src.comum.formatacao import fmt_data
from src.comum.perfil import arquivo_perfil, carregar_perfil

IGNORADOS = {"etapa.json"}
CODIGO_SEM_DIFERENCA = 0
CODIGO_ERRO = 1
CODIGO_DIFERENCAS = 6
MAXIMO_CELULAS_POR_ABA = 20  # células listadas por aba; o total aparece numa linha à parte
TOLERANCIA_RELATIVA_CELULA = 1e-12  # células numéricas: diferença abaixo disso é resto de ponto flutuante


def _arquivos(pasta: Path, ignorar: Iterable[str] = ()) -> Dict[str, Path]:
    """Arquivos da pasta por caminho relativo, sem o ``etapa.json`` e sem os caminhos de ``ignorar`` (arquivo ou pasta)."""
    fora = tuple(ignorar)
    arquivos = {}
    for p in sorted(pasta.rglob("*")):
        relativo = p.relative_to(pasta).as_posix()
        if p.is_file() and p.name not in IGNORADOS and not any(relativo == f or relativo.startswith(f"{f}/")
                                                                 for f in fora):
            arquivos[relativo] = p
    return arquivos


def _iguais(a: object, b: object) -> bool:
    if isinstance(a, float) and isinstance(b, float) and a != a and b != b:  # NaN nas duas
        return True
    numeros = (int, float)
    if isinstance(a, numeros) and isinstance(b, numeros) and not isinstance(a, bool) and not isinstance(b, bool):
        return math.isclose(a, b, rel_tol=TOLERANCIA_RELATIVA_CELULA, abs_tol=0.0)
    return a == b


def _comparar_planilhas(atual: Path, referencia: Path, nome: str) -> List[str]:
    """Diferenças entre duas planilhas: ordem das abas, quantidade de linhas e cada célula."""
    wa = load_workbook(atual, read_only=True, data_only=True)
    wr = load_workbook(referencia, read_only=True, data_only=True)
    try:
        if wa.sheetnames != wr.sheetnames:
            so_atual = [a for a in wa.sheetnames if a not in wr.sheetnames]
            so_ref = [a for a in wr.sheetnames if a not in wa.sheetnames]
            return [f"{nome}: abas diferentes ou em outra ordem (só na atual: {so_atual}; só na referência: {so_ref})"]
        diferencas: List[str] = []
        for aba in wa.sheetnames:
            linhas_a = list(wa[aba].iter_rows(values_only=True))
            linhas_r = list(wr[aba].iter_rows(values_only=True))
            if len(linhas_a) != len(linhas_r):
                diferencas.append(f"{nome} [{aba}]: {len(linhas_a)} linhas na atual e {len(linhas_r)} na referência")
            celulas: List[str] = []
            for i, (la, lr) in enumerate(zip(linhas_a, linhas_r), start=1):
                n = max(len(la), len(lr))
                la = tuple(la) + (None,) * (n - len(la))
                lr = tuple(lr) + (None,) * (n - len(lr))
                for j, (va, vr) in enumerate(zip(la, lr), start=1):
                    if not _iguais(va, vr):
                        celulas.append(f"{nome} [{aba}] {get_column_letter(j)}{i}: {va!r} na atual, {vr!r} na referência")
            diferencas += celulas[:MAXIMO_CELULAS_POR_ABA]
            if len(celulas) > MAXIMO_CELULAS_POR_ABA:
                diferencas.append(f"{nome} [{aba}]: mais {len(celulas) - MAXIMO_CELULAS_POR_ABA} células diferentes "
                                  f"({len(celulas)} no total)")
        return diferencas
    finally:
        wa.close()
        wr.close()


def comparar(pasta_atual: Path, pasta_referencia: Path, ignorar_na_referencia: Iterable[str] = ()) -> List[str]:
    """Lista as diferenças entre o relatório de ``pasta_atual`` e o de ``pasta_referencia`` (vazia se não houver).

    ``ignorar_na_referencia``: arquivos ou pastas da referência que não fazem parte do relatório.
    """
    for pasta in (pasta_atual, pasta_referencia):
        if not Path(pasta).is_dir():
            raise FileNotFoundError(f"pasta inexistente: {pasta}")
    atual = _arquivos(Path(pasta_atual))
    referencia = _arquivos(Path(pasta_referencia), ignorar_na_referencia)
    diferencas: List[str] = []
    for nome in sorted(set(atual) | set(referencia)):
        if nome not in atual:
            diferencas.append(f"{nome}: só na referência")
        elif nome not in referencia:
            diferencas.append(f"{nome}: só na pasta atual")
        elif nome.lower().endswith(".xlsx"):
            diferencas += _comparar_planilhas(atual[nome], referencia[nome], nome)
        elif atual[nome].read_bytes() != referencia[nome].read_bytes():
            diferencas.append(f"{nome}: conteúdo diferente ({atual[nome].stat().st_size} bytes na atual, "
                              f"{referencia[nome].stat().st_size} na referência)")
    return diferencas


def executar_comparacao(pasta_atual: Path, pasta_referencia: Path) -> int:
    """Compara, lista as diferenças na saída padrão e devolve o código (0, 6 ou 1)."""
    try:
        diferencas = comparar(Path(pasta_atual), Path(pasta_referencia))
    except Exception as exc:  # inclui planilha ilegível ou corrompida
        print(f"Erro na comparação: {exc}", file=sys.stderr)
        return CODIGO_ERRO
    for d in diferencas:
        print(d)
    print(f"{len(diferencas)} diferença(s)." if diferencas else "Nenhuma diferença.")
    return CODIGO_DIFERENCAS if diferencas else CODIGO_SEM_DIFERENCA


# ---------------------------------------------------------------------------
# comparar --todas [--coleta] (spec 006, decisão R22; contrato cli-referencias.md)
# ---------------------------------------------------------------------------

# Na pasta da referência, o que não é relatório
FORA_DO_RELATORIO = ("referencia.json", "perfil.toml", "coleta")
ETAPAS_REFEITAS = ("tratamento", "conferencia", "analises", "relatorio")
# Coleta refeita: registros só informativos (sem o código 6) e colunas que mudam a cada execução
COLETA_INFORMATIVOS = {"dicionarios.csv", "datas_obtencao.csv", "dicionario_evt.json"}
COLUNAS_EXECUCAO = {"data_hora_processamento"}
# Colunas de tempo das séries extraídas, na ordem de preferência: coluna, formato e duração do intervalo da linha
COLUNAS_TEMPO = (
    ("din_instante", None, pd.Timedelta(hours=1)),
    ("dia", None, pd.Timedelta(days=1)),
    ("dat_mesreferencia", None, pd.DateOffset(months=1)),
    ("din_mes", None, pd.DateOffset(months=1)),
    ("dat_periodo", "%m/%Y", pd.DateOffset(months=1)),
    ("din_ano", "%Y", pd.DateOffset(years=1)),
)


def _series_da_coleta() -> Dict[str, str]:
    """Série extraída pela Coleta -> auditoria do mesmo conjunto (``evt_extraido.csv`` -> ``auditoria_evt.csv``)."""
    arq = caminhos.ARQUIVOS_COLETA
    return {arq[chave]: arq[f"auditoria_{chave}"] for chave in arq if f"auditoria_{chave}" in arq}


def _ler_tabela(caminho: Path) -> pd.DataFrame:
    if caminho.suffix == ".parquet":
        return pd.read_parquet(caminho)
    return pd.read_csv(caminho, sep=";", dtype=str, keep_default_na=False)


def _no_periodo(tabela: pd.DataFrame, inicio: pd.Timestamp, fim: pd.Timestamp) -> pd.DataFrame:
    """Linhas cujo intervalo (hora, dia, mês ou ano da linha) toca o período; as linhas sem data ficam."""
    comeco = pd.Series(pd.NaT, index=tabela.index, dtype="datetime64[ns]")
    final = comeco.copy()
    for coluna, formato, duracao in COLUNAS_TEMPO:
        if coluna not in tabela.columns:
            continue
        valores = tabela[coluna]
        if not pd.api.types.is_datetime64_any_dtype(valores):
            valores = pd.to_datetime(valores.astype("string").replace("", pd.NA), format=formato, errors="coerce")
        valores = valores.astype("datetime64[ns]")
        novos = comeco.isna() & valores.notna()
        comeco = comeco.where(~novos, valores)
        final = final.where(~novos, valores + duracao)
    return tabela[comeco.isna() | ((comeco <= fim) & (final > inicio))]


def _linhas_diferentes(congelada: pd.DataFrame, refeita: pd.DataFrame) -> Tuple[int, Set[str]]:
    """Linhas que não coincidem (uma linha trocada conta uma vez) e os arquivos de origem dessas linhas."""
    colunas = list(dict.fromkeys([*congelada.columns, *refeita.columns]))
    a = congelada.reindex(columns=colunas).astype(str)
    b = refeita.reindex(columns=colunas).astype(str)
    a["_ordem"] = a.groupby(colunas, sort=False).cumcount()
    b["_ordem"] = b.groupby(colunas, sort=False).cumcount()
    juntas = a.merge(b, how="outer", on=[*colunas, "_ordem"], indicator=True)
    so_congelada = juntas[juntas["_merge"] == "left_only"]
    so_refeita = juntas[juntas["_merge"] == "right_only"]
    origens: Set[str] = set()
    if "arquivo_origem" in juntas.columns:
        origens = set(so_congelada["arquivo_origem"]) | set(so_refeita["arquivo_origem"])
    return max(len(so_congelada), len(so_refeita)), origens


def _contagens_diferentes(congelada: pd.DataFrame, refeita: pd.DataFrame) -> int:
    """Arquivos da Coleta congelada cujas contagens (colunas numéricas, sem as de execução) mudaram na refeita."""
    chaves = [c for c in ("conjunto", "arquivo", "nome_arquivo") if c in congelada.columns]
    if not chaves:
        return _linhas_diferentes(congelada.drop(columns=list(COLUNAS_EXECUCAO), errors="ignore"),
                                  refeita.drop(columns=list(COLUNAS_EXECUCAO), errors="ignore"))[0]
    contagens = [c for c in congelada.columns if c not in chaves and c not in COLUNAS_EXECUCAO
                 and pd.api.types.is_numeric_dtype(congelada[c])]

    def indexar(tabela: pd.DataFrame) -> pd.DataFrame:
        tabela = tabela.assign(_ordem=tabela.groupby(chaves).cumcount())
        return tabela.set_index([*chaves, "_ordem"]).reindex(columns=contagens)

    a = indexar(congelada)
    b = indexar(refeita).reindex(a.index)  # só os arquivos da referência
    iguais = (a == b) | (a.isna() & b.isna())
    return int((~iguais.all(axis=1)).sum())


def _publicacao(auditoria: Path, origens: Set[str]) -> str:
    """Publicação mais recente no ONS dos arquivos de origem, pela auditoria da Coleta refeita."""
    if not origens or not auditoria.is_file():
        return "não registrada"
    tabela = pd.read_csv(auditoria, sep=";", dtype=str, keep_default_na=False)
    chave = "nome_arquivo" if "nome_arquivo" in tabela.columns else "arquivo"
    if chave not in tabela.columns or "data_publicacao" not in tabela.columns:
        return "não registrada"
    datas = [d for d in tabela.loc[tabela[chave].isin(origens), "data_publicacao"] if d]
    return fmt_data(max(datas)) if datas else "não registrada"


def _comparar_coleta(slug: str, congelada: Path, refeita: Path, inicio: pd.Timestamp,
                     fim: pd.Timestamp) -> Tuple[List[str], List[str]]:
    """Diferenças (contam para o código 6) e informações entre a Coleta congelada e a refeita."""
    series = _series_da_coleta()

    def nomes(pasta: Path) -> Set[str]:
        return {p.name for p in pasta.iterdir() if p.is_file() and p.name not in IGNORADOS and p.suffix != ".bak"}

    nomes_congelada, nomes_refeita = nomes(congelada), nomes(refeita)
    diferencas: List[str] = []
    informacoes: List[str] = []
    for nome in sorted(nomes_congelada | nomes_refeita):
        prefixo = f"{slug}, Coleta: {nome}"
        if (nome in nomes_congelada and nome in nomes_refeita
                and (congelada / nome).read_bytes() == (refeita / nome).read_bytes()):
            continue
        if nome in COLETA_INFORMATIVOS:
            informacoes.append(f"{slug}, Coleta (informação): {nome}: diferente do congelado")
        elif nome not in nomes_refeita:
            diferencas.append(f"{prefixo}: ausente na Coleta refeita")
        elif nome not in nomes_congelada:
            diferencas.append(f"{prefixo}: só na Coleta refeita")
        elif nome.startswith("auditoria_"):
            n = _contagens_diferentes(pd.read_csv(congelada / nome, sep=";"), pd.read_csv(refeita / nome, sep=";"))
            if n:
                diferencas.append(f"{prefixo}: {n} arquivos com contagens diferentes")
        elif nome in series:
            n, origens = _linhas_diferentes(_no_periodo(_ler_tabela(congelada / nome), inicio, fim),
                                            _no_periodo(_ler_tabela(refeita / nome), inicio, fim))
            if n:
                diferencas.append(f"{prefixo}: {n} linhas diferentes no período da referência (publicação no ONS: "
                                  f"{_publicacao(refeita / series[nome], origens)})")
        elif Path(nome).suffix in (".csv", ".parquet"):
            n, _ = _linhas_diferentes(_ler_tabela(congelada / nome), _ler_tabela(refeita / nome))
            if n:
                diferencas.append(f"{prefixo}: {n} linhas diferentes")
        else:
            diferencas.append(f"{prefixo}: conteúdo diferente")
    return diferencas, informacoes


def _espaco_com_perfil(raiz: Path, pasta_referencia: Path, slug: str) -> None:
    (raiz / "usinas" / slug).mkdir(parents=True)
    shutil.copy2(pasta_referencia / referencias.NOME_PERFIL, raiz / "usinas" / slug / "perfil.toml")


def _refazer_relatorio(pasta_referencia: Path, referencia: Dict[str, Any], slug: str, raiz: Path) -> List[str]:
    """Refaz as etapas 2 a 5 em ``raiz``, a partir da Coleta e do perfil congelados, e compara com a referência."""
    _espaco_com_perfil(raiz, pasta_referencia, slug)
    shutil.copytree(pasta_referencia / referencias.PASTA_COLETA, raiz / "data" / "usinas" / slug / "coleta")
    data_geracao = datetime.strptime(referencia["data_geracao"], referencias.FORMATO_DATA_GERACAO)
    with caminhos.espaco_isolado(raiz, raw=raiz / "raw_vazio"):  # as etapas 2 a 5 não leem data/raw/
        perfil = carregar_perfil(slug)
        for etapa in ETAPAS_REFEITAS:
            opcoes = {"data_geracao": data_geracao} if etapa == "relatorio" else {}
            codigo = pipeline.executar_etapa(etapa, perfil, **opcoes)
            if codigo not in pipeline.CODIGOS_CONCLUIDA:
                raise RuntimeError(f"a etapa '{etapa}' refeita terminou com o código {codigo}")
        refeito = caminhos.pasta_relatorio(slug)
    return comparar(refeito, pasta_referencia, FORA_DO_RELATORIO)


def _refazer_coleta(pasta_referencia: Path, referencia: Dict[str, Any], slug: str, raiz: Path,
                    raw: Path) -> Tuple[List[str], List[str]]:
    """Refaz a Coleta com ``--sem-portal`` e o perfil congelado em ``raiz`` e compara com a Coleta congelada."""
    periodo = referencia.get("periodo") or {}
    if not periodo.get("inicio") or not periodo.get("fim"):
        raise RuntimeError("referencia.json sem o período da série de referência")
    _espaco_com_perfil(raiz, pasta_referencia, slug)
    with caminhos.espaco_isolado(raiz, raw=raw):  # lê os brutos do projeto; com --sem-portal, não grava neles
        codigo = pipeline.executar_etapa("coleta", carregar_perfil(slug), sem_portal=True)
        if codigo != pipeline.CODIGO_SUCESSO:
            raise RuntimeError(f"a Coleta refeita terminou com o código {codigo}")
        refeita = caminhos.pasta_etapa(slug, "coleta")
    return _comparar_coleta(slug, pasta_referencia / referencias.PASTA_COLETA, refeita,
                            pd.Timestamp(periodo["inicio"]), pd.Timestamp(periodo["fim"]))


def _comparar_referencia(pasta: Path, coleta: bool, raw: Path) -> Tuple[List[str], List[str], Optional[str]]:
    """Diferenças, informações e o aviso de perfil alterado de uma referência."""
    slug = pasta.name
    referencia = referencias.conferir_congelados(pasta)
    atual = arquivo_perfil(slug)
    aviso = None
    if not atual.is_file() or referencias.sha256_arquivo(atual) != referencia["perfil"]:
        aviso = f"{slug}: aviso: o perfil atual difere do congelado na referência (a comparação usou o congelado)"
    temporaria = Path(tempfile.mkdtemp(prefix=f"comparar_{slug}_"))
    try:
        diferencas = _refazer_relatorio(pasta, referencia, slug, temporaria / "etapas")
        informacoes: List[str] = []
        if coleta:
            da_coleta, informacoes = _refazer_coleta(pasta, referencia, slug, temporaria / "coleta", raw)
            diferencas += da_coleta
    finally:
        shutil.rmtree(temporaria, ignore_errors=True)
    return diferencas, informacoes, aviso


def executar_comparacao_todas(coleta: bool = False) -> int:
    """``comparar --todas [--coleta]``: refaz e compara cada referência. Códigos 0, 1 e 6."""
    pastas = referencias.referencias_existentes()
    if not pastas:
        print(f"Nenhum relatório de referência em {caminhos.RELATORIOS_REFERENCIA_DIR}.", file=sys.stderr)
        return CODIGO_ERRO
    raw = caminhos.RAW_DATA_DIR
    com_diferenca = com_erro = 0
    for pasta in pastas:
        try:
            diferencas, informacoes, aviso = _comparar_referencia(pasta, coleta, raw)
        except Exception as exc:  # Coleta ou perfil congelados inválidos, ou erro numa etapa refeita
            print(f"{pasta.name}: erro: {exc}", file=sys.stderr)
            com_erro += 1
            continue
        if aviso:
            print(aviso)
        if diferencas:
            com_diferenca += 1
            print(f"{pasta.name}: {len(diferencas)} diferenças")
            for d in diferencas:
                print(d)
        else:
            print(f"{pasta.name}: nenhuma diferença")
        for informacao in informacoes:
            print(informacao)
    print(f"Referências comparadas: {len(pastas) - com_erro}; com diferença: {com_diferenca}.")
    if com_erro:
        print(f"Referências com erro: {com_erro}.", file=sys.stderr)
        return CODIGO_ERRO
    return CODIGO_DIFERENCAS if com_diferenca else CODIGO_SEM_DIFERENCA
