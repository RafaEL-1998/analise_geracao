"""Relatório de carteira (spec 006, fase E, US5; decisão R21; rascunho 005, FR-049; contrato cli-carteira; T132).

Sobre um catálogo fictício de GO, com resultados gravados de usinas fictícias (``tests/fixtures/carteira_ficticia``):
- recorte por estado, por tipo e pelos dois, com o nome da pasta; sem nenhuma das duas opções, código 2;
- usinas sem resultado com cada motivo; as linhas de conjunto não contam;
- ordenação por possíveis problemas, pontos de atenção e nome, com a usina sem regra avaliada no fim;
- panorama só com o estado, com o Tipo III como previsão e o grupo sem valor como "sem valor";
- planilha com as seis abas; indicadores vazios, e não zero; nenhum termo de parecer;
- nenhuma etapa executada (os ``etapa.json`` iguais); código 5 sem catálogo ou sem usina com resultado.
"""

from __future__ import annotations

import calendar
import logging
import re
import shutil
from pathlib import Path
from typing import Dict

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.__main__ import main
from src.comum import caminhos
from tests.fixtures import brutos_ficticios as bf
from tests.fixtures import carteira_ficticia as cf

PERFIS = Path(__file__).resolve().parents[1] / "fixtures" / "usinas_ficticias"
DATA = "22/10/2026 18:00"
ABAS = ["USINAS", "INDICADORES", "ORDENACAO", "SEM_RESULTADO", "PANORAMA", "FONTES"]
PARECER = re.compile(r"satisfat|cumpre", re.IGNORECASE)
TIPO_III = "Pequenas Usinas (Tipo III)"


def _perfil(raiz: Path, origem: str, slug: str) -> None:
    (raiz / "usinas" / slug).mkdir(parents=True)
    shutil.copy(PERFIS / origem / "perfil.toml", raiz / "usinas" / slug / "perfil.toml")


def _horas(mes: str) -> int:
    ano, m = map(int, mes.split("-"))
    return calendar.monthrange(ano, m)[1] * 24


def montar_projeto(raiz: Path) -> None:
    """Catálogo de GO (e uma usina de MS), perfis e resultados gravados das Análises."""
    f = bf.FICTICIAS
    for tipo in ("ute", "ufv", "pch", "eol"):
        _perfil(raiz, tipo, f"{tipo}_ficticia")
    extras = {  # slug: (CEG, nome, trocas no texto do perfil)
        "ute_alfa": ("UTE.GN.GO.000201-0.01", "UTE Alfa", {}),
        "ute_beta": ("UTE.GN.GO.000202-0.01", "UTE Beta", {}),
        "ute_rascunho": ("UTE.GN.GO.000203-0.01", "UTE Rascunho",
                         {'modalidade = "TIPO I"': 'modalidade = "TIPO I"\nsituacao = "rascunho"'}),
        "ute_invalida": ("UTE.GN.GO.000204-0.01", "UTE Inválida", {'estado = "GO"': 'estado = "G0"'}),
        "ute_antiga": ("UTE.GN.GO.000205-0.01", "UTE Antiga", {}),
        "ute_desatualizada": ("UTE.GN.GO.000206-0.01", "UTE Desatualizada", {}),
        "ute_falha": ("UTE.GN.GO.000207-0.01", "UTE Falha", {}),
        "ute_sem_campo": ("UTE.GN.GO.000208-0.01", "UTE Sem Campo", {}),
        "dup_a": ("UTE.GN.GO.000299-0.01", "UTE Duplicada A", {}),
        "dup_b": ("UTE.GN.GO.000299-0.01", "UTE Duplicada B", {}),
    }
    for slug, (ceg, nome, trocas) in extras.items():
        texto = (PERFIS / "ute" / "perfil.toml").read_text(encoding="utf-8")
        texto = texto.replace('slug = "ute_ficticia"', f'slug = "{slug}"', 1)
        texto = texto.replace(f'ceg = "{f["ute"]["ceg"]}"', f'ceg = "{ceg}"', 1)
        texto = texto.replace('nome = "UTE Fictícia"', f'nome = "{nome}"', 1)
        for antigo, novo in trocas.items():
            texto = texto.replace(antigo, novo, 1)
        (raiz / "usinas" / slug).mkdir(parents=True)
        (raiz / "usinas" / slug / "perfil.toml").write_text(texto, encoding="utf-8")

    usinas = [
        cf.linha_usina(f["ute"]["ceg"], "UTE FICTICIA", "UTE", "GO", "proprio"),
        cf.linha_usina(f["ufv"]["ceg"], "UFV FICTICIA", "UFV", "GO", "proprio", "TIPO II-C"),
        cf.linha_usina(f["pch"]["ceg"], "PCH FICTICIA", "PCH", "GO", "conjunto", "TIPO II-C"),
        cf.linha_usina(f["eol"]["ceg"], "EOL FICTICIA", "EOL", "GO", "proprio"),
        *[cf.linha_usina(ceg, nome.upper(), "UTE", "GO", "proprio") for ceg, nome, _ in extras.values()
          if not ceg.endswith("000299-0.01")],
        cf.linha_usina("UTE.GN.GO.000299-0.01", "UTE DUPLICADA", "UTE", "GO", "proprio"),
        cf.linha_usina("CGH.PH.GO.000301-0.01", "CGH AGREGADA", "CGH", "GO", "agregado", "TIPO III"),
        cf.linha_usina("UHE.PH.GO.000302-0.01", "UHE SEM PERFIL", "UHE", "GO", "ausente"),
        {**cf.linha_usina("-", "CONJ. FICTICIO", "conjunto", "GO", "", "TIPO II-C"), "ceg": "CJU_GOCJFI"},
        cf.linha_usina("UFV.RS.MS.000303-0.01", "UFV DE MS", "UFV", "MS", "ausente", "TIPO II-C"),
    ]
    agregados = []
    for i, mes in enumerate(cf.MESES):
        agregados.append(cf.linha_agregado("GO", mes, "Pequenas Usinas (MMGD)", "FOTOVOLTAICA", "PQU GOGO FOTOV MMGD",
                                           1000.0 + 10 * i, _horas(mes), False))
        agregados.append(cf.linha_agregado("GO", mes, TIPO_III, "TÉRMICA", "PQU GOGO TERM", 300.0, _horas(mes), True))
        sem_valor = i >= 7  # o grupo publicado sem nenhum valor nos cinco últimos meses
        agregados.append(cf.linha_agregado("GO", mes, "TIPO III", "TÉRMICA", "UTE TIPO III GO", 0.0,
                                           0 if sem_valor else _horas(mes), True))
        agregados.append(cf.linha_agregado("MS", mes, "Pequenas Usinas (MMGD)", "FOTOVOLTAICA", "PQU MSMS FOTOV MMGD",
                                           999.0, _horas(mes), False))
    cf.gravar_catalogo(raiz, usinas, agregados)

    cf.gravar_analises(raiz, "ute_ficticia", cf.indicadores(
        "UTE", contagens={"possiveis_problemas": 2, "confirmar_agente": 7}, disponibilidade=20.9, teifa=30.5,
        atendimento_despacho=6.25))
    cf.gravar_analises(raiz, "ufv_ficticia", cf.indicadores(
        "UFV", ano=2026, parcial=True, contagens={"possiveis_problemas": 2, "pontos_atencao": 3},
        fator_capacidade=17.2, energia_cortada=16.3))
    cf.gravar_analises(raiz, "pch_ficticia", cf.indicadores("PCH", avaliadas=0, nao_avaliadas=11),
                       nivel_geracao="conjunto")
    cf.gravar_analises(raiz, "ute_alfa", cf.indicadores("UTE", contagens={"pontos_atencao": 5}))
    cf.gravar_analises(raiz, "ute_beta", cf.indicadores("UTE", contagens={"pontos_atencao": 5}))
    cf.gravar_analises(raiz, "ute_antiga", cf.indicadores("UTE"), versao_formato=2)
    cf.gravar_analises(raiz, "ute_desatualizada", cf.indicadores("UTE"), status="desatualizada")
    cf.gravar_analises(raiz, "ute_falha", cf.indicadores("UTE"), status="falha")
    cf.gravar_analises(raiz, "ute_sem_campo", {})


def _etapas(raiz: Path) -> Dict[str, bytes]:
    return {str(p.relative_to(raiz)): p.read_bytes() for p in sorted((raiz / "data" / "usinas").rglob("etapa.json"))}


@pytest.fixture(scope="module")
def projeto(tmp_path_factory):
    raiz = tmp_path_factory.mktemp("carteira")
    montar_projeto(raiz)
    antes = _etapas(raiz)
    with caminhos.espaco_isolado(raiz):
        codigos = {nome: main(["carteira", *opcoes, "--data-geracao", DATA])
                   for nome, opcoes in (("go", ["--estado", "GO"]), ("brasil_ute", ["--tipo", "UTE"]),
                                        ("go_ufv", ["--estado", "go", "--tipo", "ufv"]))}
    return {"raiz": raiz, "codigos": codigos, "etapas_antes": antes}


def _pasta(projeto, nome: str) -> Path:
    return projeto["raiz"] / "reports" / "carteiras" / nome


def _md(projeto, nome: str) -> str:
    return (_pasta(projeto, nome) / "relatorio_carteira.md").read_text(encoding="utf-8")


def _aba(projeto, nome: str, aba: str) -> pd.DataFrame:
    return pd.read_excel(_pasta(projeto, nome) / "carteira.xlsx", sheet_name=aba)


def test_recorte_por_estado_por_tipo_e_pelos_dois(projeto) -> None:
    assert projeto["codigos"] == {"go": 0, "brasil_ute": 0, "go_ufv": 0}
    for nome in ("go", "brasil_ute", "go_ufv"):
        pasta = _pasta(projeto, nome)
        assert {p.name for p in pasta.iterdir()} == {"relatorio_carteira.pdf", "relatorio_carteira.md", "carteira.xlsx",
                                                    "figures"}
        assert (pasta / "relatorio_carteira.pdf").read_bytes()[:5] == b"%PDF-"
    assert set(_aba(projeto, "go_ufv", "USINAS")["slug"]) == {"ufv_ficticia"}
    assert set(_aba(projeto, "brasil_ute", "USINAS")["slug"]) == {"ute_ficticia", "ute_alfa", "ute_beta"}
    assert "UFV DE MS" not in set(_aba(projeto, "go", "SEM_RESULTADO")["nome"])


@pytest.mark.parametrize("opcoes", [["carteira"], ["carteira", "--tipo", "conjunto"], ["carteira", "--estado", "GOI"]])
def test_opcao_invalida_com_codigo_2(opcoes) -> None:
    with pytest.raises(SystemExit) as saida:
        main(opcoes)
    assert saida.value.code == 2


def test_sem_resultado_com_cada_motivo_e_sem_as_linhas_de_conjunto(projeto) -> None:
    sem = _aba(projeto, "go", "SEM_RESULTADO").fillna("")
    motivos = dict(zip(sem["nome"], sem["motivo"]))
    assert motivos == {
        "CGH AGREGADA": "só agregado no ONS (ver o panorama)",
        "UHE SEM PERFIL": "sem perfil",
        "UTE RASCUNHO": "perfil em rascunho",
        "UTE INVÁLIDA": "perfil inválido",
        "UTE DUPLICADA": "perfil inválido",
        "EOL FICTICIA": "Análises ausentes",
        "UTE ANTIGA": "Análises em formato antigo",
        "UTE SEM CAMPO": "Análises em formato antigo",
        "UTE FALHA": "Análises com falha",
        "UTE DESATUALIZADA": "Análises desatualizadas",
    }
    detalhes = dict(zip(sem["nome"], sem["detalhe"]))
    assert detalhes["UTE DUPLICADA"] == "dois perfis com o mesmo CEG: dup_a e dup_b"
    assert detalhes["UTE SEM CAMPO"] == "ute_sem_campo; sem os indicadores da carteira"
    assert list(sem["motivo"]) == sorted(sem["motivo"], key=[
        "só agregado no ONS (ver o panorama)", "sem perfil", "perfil em rascunho", "perfil inválido",
        "Análises ausentes", "Análises em formato antigo", "Análises com falha", "Análises desatualizadas"].index)
    assert "CONJ. FICTICIO" not in set(sem["nome"]) | set(_aba(projeto, "go", "USINAS")["nome"])
    md = _md(projeto, "go")
    assert "15 usinas no recorte: 5 com resultado e 10 sem resultado." in md
    assert "A linha de conjunto de usinas do catálogo não conta como usina." in md


def test_ordenacao_por_problemas_atencao_e_nome(projeto) -> None:
    ordem = _aba(projeto, "go", "ORDENACAO")
    assert list(ordem["slug"]) == ["ufv_ficticia", "ute_ficticia", "ute_alfa", "ute_beta", "pch_ficticia"]
    assert list(ordem["posicao"].fillna(0).astype(int)) == [1, 2, 3, 4, 0]  # sem regra avaliada: sem posição
    pch = ordem[ordem["slug"] == "pch_ficticia"].iloc[0]
    assert pd.isna(pch["possiveis_problemas"]) and int(pch["regras_avaliadas"]) == 0  # vazio, e não zero
    md = _md(projeto, "go")
    assert md.index("| 1 | UFV Fictícia |") < md.index("| 2 | UTE Fictícia |") < md.index("| — | PCH Fictícia |")
    assert ("A conclusão de PCH Fictícia não tem base no nível da usina: nenhuma regra foi avaliada com dado da usina"
            in md)
    assert (_pasta(projeto, "go") / "figures" / "01_itens_da_conclusao_por_usina.png").is_file()


def test_indicadores_vazios_e_nao_zero(projeto) -> None:
    ind = _aba(projeto, "go", "INDICADORES")
    ute = ind[ind["slug"] == "ute_ficticia"].set_index("indicador")
    assert ute.loc["disponibilidade", "valor"] == pytest.approx(20.9)
    assert pd.isna(ute.loc["fator_capacidade", "valor"]) and pd.isna(ute.loc["energia_cortada", "valor"])
    assert ute.loc["atendimento_despacho", "detalhe"] == "2 de 32 horas com programação"
    md = _md(projeto, "go")
    assert "| UTE Fictícia | 2025 | — | 20,9% | 30,5% | — | 6,2% (2 de 32 h) |" in md
    assert "| UFV Fictícia | 2026\\* | 17,2% | — | — | 16,3% | — |" in md
    # na carteira de térmicas, só as colunas que se aplicam ao tipo
    assert "| Usina | Ano | Disponibilidade | TEIFa | Atendimento ao despacho |" in _md(projeto, "brasil_ute")


def test_panorama_so_com_estado_e_tipo_iii_como_previsao(projeto) -> None:
    panorama = _aba(projeto, "go", "PANORAMA")
    assert set(panorama["id_estado"]) == {"GO"}
    assert set(panorama.loc[panorama["previsao"], "modalidade"]) == {TIPO_III, "TIPO III"}
    assert "Pequenas usinas Tipo III térmica (previsão)" in set(panorama["serie"])
    sem_valor = panorama[panorama["grupo"] == "UTE TIPO III GO"]
    assert sem_valor["energia_mwh"].isna().sum() == 5  # nenhuma hora com valor: sem energia, nem zero
    md = _md(projeto, "go")
    assert "## 5. Panorama dos agregados de GO" in md
    assert "| set/2026 | 1.110 | 300 | sem valor |" in md
    assert "Horas sem valor: 3.672 em UTE TIPO III GO (de mai/2026 a set/2026)." in md
    assert (_pasta(projeto, "go") / "figures" / "02_panorama_agregados_mensal.png").is_file()
    # por estado e tipo, só o tipo do ONS correspondente; só por tipo, omitido com o motivo
    assert set(_aba(projeto, "go_ufv", "PANORAMA")["tipo_ons"]) == {"FOTOVOLTAICA"}
    assert list(_aba(projeto, "brasil_ute", "PANORAMA").columns) == ["motivo"]
    md_ute = _md(projeto, "brasil_ute")
    assert "Panorama dos agregados omitido: o ONS publica os agregados por área do estado" in md_ute
    assert "## 5. Panorama" not in md_ute
    assert not (_pasta(projeto, "brasil_ute") / "figures" / "02_panorama_agregados_mensal.png").exists()


def test_planilha_com_as_seis_abas(projeto) -> None:
    for nome in ("go", "brasil_ute", "go_ufv"):
        livro = load_workbook(_pasta(projeto, nome) / "carteira.xlsx", read_only=True)
        try:
            assert livro.sheetnames == ABAS
        finally:
            livro.close()


LEGENDAS = ("Fonte dos dados:", "Calculado neste relatório a partir de:")


def _depois(linhas, i: int) -> str:
    return next((linha for linha in linhas[i + 1:] if linha.strip()), "")


def test_legenda_de_fonte_abaixo_de_cada_tabela_e_figura(projeto) -> None:
    """T163 (princípio VII): a carteira segue a regra de legendas do relatório da usina."""
    for nome in ("go", "brasil_ute"):
        linhas = _md(projeto, nome).splitlines()
        fins = [i for i, linha in enumerate(linhas) if linha.startswith("|") and not _depois(linhas, i).startswith("|")]
        figuras = [i for i, linha in enumerate(linhas) if linha.startswith("*Figura ")]
        assert fins and all(_depois(linhas, i).startswith(LEGENDAS) for i in fins), nome
        assert all(_depois(linhas, i).startswith(LEGENDAS) for i in figuras), nome
    md = _md(projeto, "go")
    assert "Modalidade das usinas (CEG de cada usina), publicado em" in md
    assert "Geração por usina (agregados de GO por modalidade, nível agregado)" in md


def test_aba_fontes_com_as_colunas_da_usina(projeto) -> None:
    """T164 (R21): a aba FONTES da carteira tem as colunas da aba FONTES da usina, uma linha por aba."""
    from src.relatorio.fontes import COLUNAS_FONTES

    fontes = _aba(projeto, "go", "FONTES")
    assert list(fontes.columns) == COLUNAS_FONTES
    assert list(fontes["aba"]) == ABAS[:-1]
    assert set(fontes["calculado_no_relatorio"]) <= {"sim", "não"}


def test_nenhum_termo_de_parecer(projeto) -> None:
    for nome in ("go", "brasil_ute", "go_ufv"):
        md = _md(projeto, nome)
        assert not PARECER.search(md)
        assert "A carteira não atribui nota nem parecer" in md


def test_nenhuma_etapa_executada(projeto) -> None:
    assert _etapas(projeto["raiz"]) == projeto["etapas_antes"]
    assert not (projeto["raiz"] / "reports" / "ute_ficticia").exists()


def test_pdf_reproduzivel_e_mensagem_do_fim(projeto, caplog) -> None:
    pdf = _pasta(projeto, "go") / "relatorio_carteira.pdf"
    antes = pdf.read_bytes()
    with caminhos.espaco_isolado(projeto["raiz"]), caplog.at_level(logging.INFO):
        assert main(["carteira", "--estado", "GO", "--data-geracao", DATA]) == 0
    assert pdf.read_bytes() == antes
    assert ("Carteira go: 5 usinas com resultado e 10 sem resultado. Pasta: reports/carteiras/go/"
            in caplog.text)
    assert not list(_pasta(projeto, "go").rglob("*.bak")) and not (_pasta(projeto, "go") / ".gravando").exists()


def test_codigo_5_sem_catalogo_ou_sem_usina_com_resultado(tmp_path, caplog) -> None:
    montar_projeto(tmp_path)
    with caminhos.espaco_isolado(tmp_path), caplog.at_level(logging.ERROR):
        assert main(["carteira", "--tipo", "CGH"]) == 5
        assert "Nenhuma usina do recorte CGH no Brasil tem as Análises concluídas" in caplog.text
        shutil.rmtree(tmp_path / "data" / "catalogo")
        assert main(["carteira", "--estado", "GO"]) == 5
        assert "O catálogo de usinas não foi montado. Execute antes: python -m src usinas" in caplog.text
    assert not (tmp_path / "reports" / "carteiras").exists()
