"""Testes unitários e de integração para o módulo processor.py."""

from __future__ import annotations

from pathlib import Path
import openpyxl
import pandas as pd
import pyarrow.parquet as pq
import pytest

from src.config import OPERATIONAL_METRIC_COLUMNS
from src.processor import (
    executar_pipeline_tratamento,
    exportar_csv,
    exportar_excel,
    exportar_parquet,
    padronizar_tipagem_numerica,
)
from tests.conftest import gerar_df_sintetico


@pytest.fixture
def amostra_df_misto() -> pd.DataFrame:
    """DataFrame com formatos mistos (vírgula regional, espaços, vazio)."""
    return pd.DataFrame(
        [
            {
                "id_subsistema": "SE",
                "nom_subsistema": "SUDESTE",
                "nom_bacia": "PARANA",
                "nom_rio": "VERDE",
                "nom_agente": "AXIA SUL",
                "nom_reservatorio": "SAO DOMINGOS",
                "cod_usina": "153",
                "din_instante": "2024-01-01 00:00:00",
                "val_geracao": "30,500",  # vírgula regional
                "val_disponibilidade": 46.8,
                "val_vazaoturbinada": " 100.0 ",  # espaços
                "val_vazaovertida": 6.0,
                "val_vazaovertidanaoturbinavel": "5",
                "val_produtividade": 0.3053061224489796,
                "val_folgadegeracao": 16.3,
                "val_energiavertida": "",  # ausente
                "val_vazaovertidaturbinavel": 1.0,
                "val_energiavertidaturbinavel": 0.3053061224489796,
                "arquivo_origem": "teste_2024.csv",
                "tipo_match": "CODIGO_E_NOME",
            }
        ]
    )


def test_padronizar_tipagem_numerica(amostra_df_misto: pd.DataFrame) -> None:
    """Converte para float64 sem transformar ausentes em zero."""
    df_tratado = padronizar_tipagem_numerica(amostra_df_misto)

    for col in OPERATIONAL_METRIC_COLUMNS:
        assert df_tratado[col].dtype == "float64", f"Coluna {col} deveria ser float64"

    assert df_tratado["val_geracao"].iloc[0] == 30.5
    assert df_tratado["val_vazaoturbinada"].iloc[0] == 100.0
    assert pd.isna(df_tratado["val_energiavertida"].iloc[0]), "Valor ausente não pode virar 0"
    assert pd.api.types.is_datetime64_any_dtype(df_tratado["din_instante"])
    assert df_tratado["cod_usina"].iloc[0] == 153


def test_padronizar_rejeita_instante_invalido(amostra_df_misto: pd.DataFrame) -> None:
    amostra_df_misto.loc[0, "din_instante"] = "data inválida"
    with pytest.raises(ValueError):
        padronizar_tipagem_numerica(amostra_df_misto)


def test_exportar_excel_celulas_numericas(tmp_path: Path, amostra_df_misto: pd.DataFrame) -> None:
    """A exportação Excel grava células numéricas nativas."""
    df_tratado = padronizar_tipagem_numerica(amostra_df_misto)
    caminho_xlsx = tmp_path / "teste_saida.xlsx"
    exportar_excel(df_tratado, caminho_xlsx)

    ws = openpyxl.load_workbook(caminho_xlsx)["UHE_SAO_DOMINGOS"]
    header = [cell.value for cell in ws[1]]
    celula = ws.cell(row=2, column=header.index("val_geracao") + 1)
    assert celula.data_type == "n"
    assert celula.value == 30.5


def test_exportar_parquet_schema(tmp_path: Path, amostra_df_misto: pd.DataFrame) -> None:
    """A exportação Parquet usa DOUBLE para as métricas."""
    df_tratado = padronizar_tipagem_numerica(amostra_df_misto)
    caminho_parquet = tmp_path / "teste_saida.parquet"
    exportar_parquet(df_tratado, caminho_parquet)

    schema = pq.read_table(caminho_parquet).schema
    for col in OPERATIONAL_METRIC_COLUMNS:
        assert schema.field(col).type == "double"


def test_exportar_csv_precisao_integral(tmp_path: Path, amostra_df_misto: pd.DataFrame) -> None:
    """O CSV mantém ';', ponto decimal e todas as casas (sem arredondar para 4 dígitos)."""
    df_tratado = padronizar_tipagem_numerica(amostra_df_misto)
    caminho_csv = tmp_path / "teste_saida.csv"
    exportar_csv(df_tratado, caminho_csv)

    relido = pd.read_csv(caminho_csv, sep=";")
    assert relido["val_produtividade"].iloc[0] == pytest.approx(0.3053061224489796, abs=1e-15)
    assert pd.isna(relido["val_energiavertida"].iloc[0])


def test_executar_pipeline_tratamento_sinaliza_anomalias(tmp_path: Path) -> None:
    """Pipeline completo sobre base sintética: exporta os formatos e a coluna de qualidade."""
    entrada = tmp_path / "consolidado.csv"
    df = gerar_df_sintetico()
    df["din_instante"] = df["din_instante"].dt.strftime("%Y-%m-%d %H:%M:%S")
    df.to_csv(entrada, sep=";", index=False)

    codigo = executar_pipeline_tratamento(input_file=entrada, output_dir=tmp_path)

    assert codigo == 0
    tratado = pd.read_parquet(tmp_path / "uhe_sao_domingos_energia_vertida_tratado.parquet")
    assert len(tratado) == len(df)
    assert "qualidade_registro" in tratado.columns
    anomalos = tratado[tratado["qualidade_registro"] != "OK"]
    assert list(anomalos["din_instante"].astype(str)) == ["2024-02-20 10:00:00"]
    assert (tmp_path / "relatorio_validacao_fisica.md").exists()
    assert (tmp_path / "uhe_sao_domingos_energia_vertida_tratado.xlsx").exists()
    assert (tmp_path / "uhe_sao_domingos_energia_vertida_tratado.csv").exists()
