"""Execução de ponta a ponta da usina fictícia e extração dos invariantes do relatório (decisão R29, spec 006).

Usado pelo gerador da fotografia (``tests/fixtures/gerar_invariantes_uhe.py``) e pelos testes de invariantes
(``tests/relatorio/test_invariantes_uhe.py``). Roda sem rede, numa raiz temporária, com os brutos sintéticos de
``tests/fixtures/brutos_ficticios.py``.
"""

from __future__ import annotations

import re
import shutil
from contextlib import ExitStack
from pathlib import Path
from typing import Any, Dict, List
from unittest import mock

import pandas as pd
import pyarrow.parquet as pq
from openpyxl import load_workbook

from src.__main__ import main
from src.analises.resultados import carregar_resultados
from src.comum import caminhos
from src.comum import perfil as modulo_perfil
from tests.fixtures.brutos_ficticios import gerar_brutos

PERFIL_FICTICIO = Path(__file__).resolve().parent / "usina_ficticia" / "perfil.toml"
SLUG = "usina_ficticia"
DATA_GERACAO = "08/10/2026 10:00"
ARQUIVO_INVARIANTES = Path(__file__).resolve().parent / "invariantes_uhe.json"
_DATA_OBTENCAO = re.compile(r"obtido em \d{2}/\d{2}/\d{4}")
_TITULOS_LEGENDA = ("Fonte dos dados:", "Calculado neste relatório a partir de:")


def executar_usina_ficticia(raiz: Path, perfil: Path = PERFIL_FICTICIO) -> Dict[str, Any]:
    """Grava o perfil e os brutos sintéticos em ``raiz`` e executa o fluxo completo com ``--sem-portal``."""
    raiz = Path(raiz)
    (raiz / "usinas" / SLUG).mkdir(parents=True, exist_ok=True)
    shutil.copy(perfil, raiz / "usinas" / SLUG / "perfil.toml")
    gerar_brutos(raiz / "data" / "raw")
    with ExitStack() as pilha:
        pilha.enter_context(mock.patch.object(caminhos, "DATA_DIR", raiz / "data"))
        pilha.enter_context(mock.patch.object(caminhos, "RAW_DATA_DIR", raiz / "data" / "raw"))
        pilha.enter_context(mock.patch.object(caminhos, "REPORTS_DIR", raiz / "reports"))
        pilha.enter_context(mock.patch.object(modulo_perfil, "RAIZ_PROJETO", raiz))
        codigo = main(["completo", "--usina", SLUG, "--sem-portal", "--data-geracao", DATA_GERACAO])
    return {"codigo": codigo, "raiz": raiz, "dados": raiz / "data" / "usinas" / SLUG,
            "relatorio": raiz / "reports" / SLUG}


def _normalizar(texto: str) -> str:
    return _DATA_OBTENCAO.sub("obtido em <data>", texto)


def _secao(linhas: List[str], titulo: str) -> List[str]:
    """Linhas da seção ``## N. <titulo>`` até a próxima seção de nível 2."""
    saida: List[str] = []
    dentro = False
    for linha in linhas:
        if linha.startswith("## "):
            dentro = bool(re.match(rf"^## \d+\. {re.escape(titulo)}$", linha))
            continue
        if dentro:
            saida.append(linha)
    return saida


def _colunas_coleta(pasta: Path) -> Dict[str, List[str]]:
    colunas: Dict[str, List[str]] = {}
    for arquivo in sorted(pasta.iterdir()):
        if not arquivo.is_file() or arquivo.name == "etapa.json" or arquivo.suffix == ".bak":
            continue
        if arquivo.suffix == ".parquet":
            colunas[arquivo.name] = list(pq.read_schema(arquivo).names)
        elif arquivo.suffix == ".csv":
            colunas[arquivo.name] = list(pd.read_csv(arquivo, sep=";", nrows=0).columns)
        else:
            colunas[arquivo.name] = []
    return colunas


def _planilha(caminho: Path) -> Dict[str, Any]:
    livro = load_workbook(caminho, read_only=True)
    try:
        abas = {}
        dicionarios: List[str] = []
        for nome in livro.sheetnames:
            linhas = livro[nome].iter_rows(values_only=True)
            cabecalho = next(linhas, ())
            abas[nome] = [None if c is None else str(c) for c in cabecalho]
            if nome == "DICIONARIOS" and "conjunto" in abas[nome]:
                indice = abas[nome].index("conjunto")
                for linha in linhas:
                    if linha[indice] not in dicionarios:
                        dicionarios.append(linha[indice])
        return {"abas": abas, "dicionarios": dicionarios}
    finally:
        livro.close()


def extrair_invariantes(execucao: Dict[str, Any]) -> Dict[str, Any]:
    """Os invariantes da decisão R29 a partir de uma execução de ``executar_usina_ficticia``."""
    from src.relatorio.conteudo import nota_conclusao

    relatorio, dados = execucao["relatorio"], execucao["dados"]
    md = _normalizar((relatorio / "relatorio_analise_estatistica.md").read_text(encoding="utf-8"))
    linhas = md.splitlines()
    inicio_sumario = next(i for i, linha in enumerate(linhas) if linha == "## Sumário")
    res = carregar_resultados(dados / "analises" / caminhos.ARQUIVOS_ANALISES["resultados"])
    planilha = _planilha(relatorio / caminhos.ARQUIVOS_RELATORIO["xlsx"])
    return {
        "md_completo": md,
        "secoes": [linha for linha in linhas if re.match(r"^## \d+\. ", linha)],
        "capa": linhas[:inicio_sumario],
        "legendas": [linha for linha in linhas if linha.startswith(_TITULOS_LEGENDA)],
        "notas": _secao(linhas, "Notas metodológicas e limitações"),
        "conclusao_md": _secao(linhas, "Conclusão"),
        "conclusao_itens": [[i["lista"], i["regra"], i["texto"]] for i in res.conclusao],
        "nota_conclusao": nota_conclusao(),
        "planilha_abas": planilha["abas"],
        "dicionarios": planilha["dicionarios"],
        "coleta": _colunas_coleta(dados / "coleta"),
    }
