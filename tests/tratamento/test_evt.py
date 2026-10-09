"""Testes do tratamento da base de EVT: tipagem numérica, exportação e validação física."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pyarrow.parquet as pq
import pytest

from src.comum.caminhos import ARQUIVOS_TRATAMENTO
from src.comum.perfil import carregar_perfil
from src.comum.regras import OPERATIONAL_METRIC_COLUMNS
from src.tratamento.evt import (
    exportar_csv,
    exportar_excel,
    exportar_parquet,
    nome_aba,
    padronizar_tipagem_numerica,
    tratar_evt,
)
from tests.conftest import gerar_df_sintetico
from tests.fixtures.brutos_ficticios import gravar_dicionario_evt

PERFIL = carregar_perfil("sao_domingos")


@pytest.fixture
def amostra_df_misto() -> pd.DataFrame:
    """DataFrame com formatos mistos (vírgula regional, espaços, vazio)."""
    return pd.DataFrame(
        [
            {
                "id_subsistema": "SE",
                "nom_subsistema": "SUDESTE",
                "nom_bacia": "PARANA",
                "nom_rio": "VERDE",
                "nom_agente": "AXIA SUL",
                "nom_reservatorio": "SAO DOMINGOS",
                "cod_usina": "153",
                "din_instante": "2024-01-01 00:00:00",
                "val_geracao": "30,500",  # vírgula regional
                "val_disponibilidade": 46.8,
                "val_vazaoturbinada": " 100.0 ",  # espaços
                "val_vazaovertida": 6.0,
                "val_vazaovertidanaoturbinavel": "5",
                "val_produtividade": 0.3053061224489796,
                "val_folgadegeracao": 16.3,
                "val_energiavertida": "",  # ausente
                "val_vazaovertidaturbinavel": 1.0,
                "val_energiavertidaturbinavel": 0.3053061224489796,
                "arquivo_origem": "teste_2024.csv",
                "tipo_match": "CODIGO_E_NOME",
            }
        ]
    )


def _arquivos(pasta: Path) -> dict:
    return {chave: pasta / ARQUIVOS_TRATAMENTO[chave]
            for chave in ("evt_parquet", "evt_csv", "evt_xlsx", "validacao_csv", "validacao_md")}


def test_nome_da_aba_vem_do_nome_da_usina() -> None:
    assert nome_aba(PERFIL) == "UHE_SAO_DOMINGOS"


def test_padronizar_tipagem_numerica(amostra_df_misto: pd.DataFrame) -> None:
    """Converte para float64 sem transformar ausentes em zero."""
    df_tratado = padronizar_tipagem_numerica(amostra_df_misto)
    for col in OPERATIONAL_METRIC_COLUMNS:
        assert df_tratado[col].dtype == "float64", f"Coluna {col} deveria ser float64"
    assert df_tratado["val_geracao"].iloc[0] == 30.5
    assert df_tratado["val_vazaoturbinada"].iloc[0] == 100.0
    assert pd.isna(df_tratado["val_energiavertida"].iloc[0]), "Valor ausente não pode virar 0"
    assert pd.api.types.is_datetime64_any_dtype(df_tratado["din_instante"])
    assert df_tratado["cod_usina"].iloc[0] == 153


def test_padronizar_rejeita_instante_invalido(amostra_df_misto: pd.DataFrame) -> None:
    amostra_df_misto.loc[0, "din_instante"] = "data inválida"
    with pytest.raises(ValueError):
        padronizar_tipagem_numerica(amostra_df_misto)


def test_exportar_excel_celulas_numericas(tmp_path: Path, amostra_df_misto: pd.DataFrame) -> None:
    """A exportação Excel grava células numéricas nativas, na aba com o nome da usina."""
    df_tratado = padronizar_tipagem_numerica(amostra_df_misto)
    caminho_xlsx = tmp_path / "teste_saida.xlsx"
    exportar_excel(df_tratado, caminho_xlsx, nome_aba(PERFIL))
    ws = openpyxl.load_workbook(caminho_xlsx)[nome_aba(PERFIL)]
    header = [cell.value for cell in ws[1]]
    celula = ws.cell(row=2, column=header.index("val_geracao") + 1)
    assert celula.data_type == "n"
    assert celula.value == 30.5


def test_exportar_parquet_schema(tmp_path: Path, amostra_df_misto: pd.DataFrame) -> None:
    """A exportação Parquet usa DOUBLE para as métricas."""
    df_tratado = padronizar_tipagem_numerica(amostra_df_misto)
    caminho_parquet = tmp_path / "teste_saida.parquet"
    exportar_parquet(df_tratado, caminho_parquet)
    schema = pq.read_table(caminho_parquet).schema
    for col in OPERATIONAL_METRIC_COLUMNS:
        assert schema.field(col).type == "double"


def test_exportar_csv_precisao_integral(tmp_path: Path, amostra_df_misto: pd.DataFrame) -> None:
    """O CSV mantém ';', ponto decimal e todas as casas (sem arredondar para 4 dígitos)."""
    df_tratado = padronizar_tipagem_numerica(amostra_df_misto)
    caminho_csv = tmp_path / "teste_saida.csv"
    exportar_csv(df_tratado, caminho_csv)
    relido = pd.read_csv(caminho_csv, sep=";")
    assert relido["val_produtividade"].iloc[0] == pytest.approx(0.3053061224489796, abs=1e-15)
    assert pd.isna(relido["val_energiavertida"].iloc[0])


def test_tratar_evt_sinaliza_anomalias(tmp_path: Path) -> None:
    """Base sintética: exporta os formatos, a validação e a coluna de qualidade."""
    entrada = tmp_path / "evt_extraido.csv"
    gravar_dicionario_evt(tmp_path / "dicionario_evt.json")  # a cópia que a Coleta grava ao lado da base
    df = gerar_df_sintetico()
    df["din_instante"] = df["din_instante"].dt.strftime("%Y-%m-%d %H:%M:%S")
    df.to_csv(entrada, sep=";", index=False)
    saida = tmp_path / "tratamento"
    codigo, gravados, resumo = tratar_evt(PERFIL, entrada, _arquivos(saida))
    assert resumo["registros"] == len(df) and resumo["sinalizados"] == 1
    assert codigo == 0 and all(p.exists() for p in gravados) and len(gravados) == 5
    tratado = pd.read_parquet(saida / ARQUIVOS_TRATAMENTO["evt_parquet"])
    assert len(tratado) == len(df) and "qualidade_registro" in tratado.columns
    anomalos = tratado[tratado["qualidade_registro"] != "OK"]
    assert list(anomalos["din_instante"].astype(str)) == ["2024-02-20 10:00:00"]


def test_tratar_evt_com_valor_negativo_e_erro(tmp_path: Path) -> None:
    """R1 (valor negativo publicado) é erro: código 1 e nada gravado."""
    entrada = tmp_path / "evt_extraido.csv"
    gravar_dicionario_evt(tmp_path / "dicionario_evt.json")
    df = gerar_df_sintetico()
    df["din_instante"] = df["din_instante"].dt.strftime("%Y-%m-%d %H:%M:%S")
    df.loc[0, "val_geracao"] = -3.0
    df.to_csv(entrada, sep=";", index=False)
    saida = tmp_path / "tratamento"
    codigo, gravados, resumo = tratar_evt(PERFIL, entrada, _arquivos(saida))
    assert resumo["violacoes"]["R1"] == 1
    assert codigo == 1 and gravados == [] and not saida.exists()


def test_nome_da_aba_sem_caracteres_proibidos_e_com_no_maximo_31() -> None:
    """Decisão de 08/10/2026 (FR-008): caracteres que o Excel não aceita viram "_", e o nome é cortado em 31."""
    from types import SimpleNamespace

    from src.tratamento.evt import CARACTERES_PROIBIDOS_ABA, nome_aba

    perfil = SimpleNamespace(usina=SimpleNamespace(nome="UHE Usina/Teste [Rio]: Ábaco do Norte Muito Comprida"))
    nome = nome_aba(perfil)
    assert len(nome) == 31 and nome.startswith("UHE_USINA_TESTE__RIO__")
    assert not any(c in nome for c in CARACTERES_PROIBIDOS_ABA)
    assert nome_aba(carregar_perfil("sao_domingos")) == "UHE_SAO_DOMINGOS"
