"""Geração do relatório: PDF, Markdown, planilha, CSV e figuras da usina, só a partir do que as etapas anteriores gravaram.

Lê ``data/usinas/<slug>/analises/resultados.pkl``, monta a tabela de parâmetros (perfil, regras gerais e datas de
obtenção) e a origem dos dados, desenha as figuras e grava ``reports/<slug>/`` (spec da Geração do relatório).
"""

from __future__ import annotations

import os
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook

from src.analises.resultados import ResultadosAnalise, carregar_resultados
from src.comum import caminhos
from src.comum.logger import setup_logger
from src.comum.regras import DEFAULT_PLOT_DPI
from src.pipeline import CODIGO_SUCESSO, ResultadoEtapa
from src.relatorio.estrutura import secoes_presentes
from src.relatorio.figuras import NOMES_FIGURAS_OPCIONAIS, gerar_graficos
from src.relatorio.fontes import origem_dos_dados
from src.relatorio.markdown import gerar_relatorio_md
from src.relatorio.parametros import (
    parametros_disponibilidade,
    parametros_geracao_cadastro,
    parametros_hidrologia,
    tabela_parametros,
)
from src.relatorio.pdf import PDFReportGenerator
from src.relatorio.planilha import exportar_tabelas

logger = setup_logger("relatorio")

# Pasta, dentro da do relatório, onde as saídas são geradas e conferidas antes de substituir as da execução anterior
PASTA_TEMPORARIA = ".gravando"


def completar_resultados(res: ResultadosAnalise) -> ResultadosAnalise:
    """Tabela de parâmetros e origem dos dados, montadas no relatório (spec das Análises, FR-042)."""
    parametros = tabela_parametros(com_indicadores=bool(res.ons), com_programacao=bool(res.programacao))
    if res.disponibilidade:
        parametros = pd.concat([parametros, parametros_disponibilidade(res)], ignore_index=True)
    if res.hidrologia:
        parametros = pd.concat([parametros, parametros_hidrologia(res)], ignore_index=True)
    if res.geracao_oficial or res.cadastro:
        parametros = pd.concat([parametros, parametros_geracao_cadastro(res)], ignore_index=True)
    res.parametros = parametros
    res.fontes = origem_dos_dados(res)
    return res


def paginas_do_pdf(caminho: Path) -> int:
    """Quantidade de páginas: objetos de página do PDF (o objeto da árvore de páginas é /Pages)."""
    return len(re.findall(rb"/Type\s*/Page(?![a-z])", Path(caminho).read_bytes()))


def _conferir_saidas(arquivos: List[Path]) -> None:
    """Conferência no disco antes da troca: cada saída existe e não está vazia; o PDF começa com %PDF-."""
    for arquivo in arquivos:
        if not arquivo.is_file() or arquivo.stat().st_size == 0:
            raise OSError(f"saída não gravada ou vazia: {arquivo.name}")
        if arquivo.suffix == ".pdf" and arquivo.read_bytes()[:5] != b"%PDF-":
            raise OSError(f"PDF inválido: {arquivo.name}")


def _trocar_saidas(temporaria: Path, pasta: Path, novos: List[Path]) -> List[Path]:
    """Substitui as saídas da execução anterior pelas novas; se algo falhar, as anteriores voltam como estavam."""
    guardadas = temporaria / ".anteriores"
    trocados: List[Tuple[Path, Optional[Path]]] = []
    try:
        for novo in novos:
            relativo = novo.relative_to(temporaria)
            destino = pasta / relativo
            destino.parent.mkdir(parents=True, exist_ok=True)
            guardada = None
            if destino.exists():
                guardada = guardadas / relativo
                guardada.parent.mkdir(parents=True, exist_ok=True)
                os.replace(destino, guardada)
            trocados.append((destino, guardada))
            os.replace(novo, destino)
    except BaseException:
        falhas = []
        for destino, guardada in reversed(trocados):  # cada arquivo volta à versão anterior, um de cada vez
            try:
                if guardada is not None:
                    if guardada.exists():
                        os.replace(guardada, destino)
                else:
                    destino.unlink(missing_ok=True)
            except OSError as exc:
                falhas.append(f"{destino.name} ({exc})")
        if falhas:  # as versões anteriores que não voltaram saem da pasta temporária, que será apagada
            preservadas = pasta / f".anteriores_{datetime.now().strftime('%Y%m%dT%H%M%S')}"
            os.replace(guardadas, preservadas)
            logger.error("Troca das saídas do relatório interrompida; não voltaram à versão anterior: %s. As versões "
                         "anteriores estão em %s.", "; ".join(falhas), preservadas)
        else:
            logger.error("Troca das saídas do relatório interrompida; as da execução anterior foram mantidas.")
        raise
    return [pasta / novo.relative_to(temporaria) for novo in novos]


def executar_relatorio(perfil: Any, data_geracao: Optional[datetime] = None) -> ResultadoEtapa:
    """Gera o relatório da usina em ``reports/<slug>/`` (código 0; falha em qualquer saída é erro, código 1).

    As saídas são geradas numa pasta temporária, conferidas e só então substituem as da execução anterior, que ficam
    como estavam se algo falhar (constituição, Requisito Técnico 3; spec da Geração do relatório, FR-003).
    """
    slug = perfil.usina.slug
    pasta = caminhos.pasta_relatorio(slug)
    arq = caminhos.ARQUIVOS_RELATORIO
    res = completar_resultados(carregar_resultados(caminhos.pasta_etapa(slug, "analises")
                                                   / caminhos.ARQUIVOS_ANALISES["resultados"]))
    temporaria = pasta / PASTA_TEMPORARIA
    if temporaria.exists():
        shutil.rmtree(temporaria)
    temporaria.mkdir(parents=True)
    try:
        figuras = gerar_graficos(res, temporaria / arq["figuras"], dpi=DEFAULT_PLOT_DPI)
        xlsx, csv = exportar_tabelas(res, temporaria / arq["xlsx"], temporaria / arq["csv"])
        gerado_em = data_geracao or datetime.now()  # a mesma data no PDF, no Markdown e no resumo
        md = gerar_relatorio_md(res, figuras, temporaria / arq["md"], data_geracao=gerado_em)
        gerador = PDFReportGenerator(res, figuras, temporaria / arq["pdf"], data_geracao=gerado_em,
                                     invariante=data_geracao is not None)
        pdf = gerador.build_pdf()
        novos: List[Path] = [pdf, md, xlsx, csv, *sorted(figuras.values())]
        _conferir_saidas(novos)
        planilha = load_workbook(xlsx, read_only=True)
        abas = len(planilha.sheetnames)
        planilha.close()  # no Windows, a planilha aberta não pode ser movida
        resumo = {
            "data_geracao": gerado_em.strftime("%d/%m/%Y %H:%M"),
            "secoes": len(secoes_presentes(res)),
            "paginas_pdf": paginas_do_pdf(pdf),
            "figuras": len(figuras),
            "abas_planilha": abas,
        }
        arquivos = _trocar_saidas(temporaria, pasta, novos)
    finally:
        shutil.rmtree(temporaria, ignore_errors=True)
    for nome in NOMES_FIGURAS_OPCIONAIS.values():  # figura opcional de uma execução anterior, não gerada agora
        antiga = pasta / arq["figuras"] / nome
        if antiga not in arquivos and antiga.exists():
            antiga.unlink()
            logger.info("Figura de execução anterior removida: %s", nome)
    return ResultadoEtapa(codigo=CODIGO_SUCESSO, arquivos=arquivos, resumo=resumo)
