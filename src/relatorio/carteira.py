"""Relatório de carteira (spec da Geração do relatório, FR-049; plano, D22; contrato
``contracts/cli-carteira.md``).

``python -m src carteira [--estado <UF>] [--tipo <tipo>] [--data-geracao ...]`` monta o relatório de uma carteira de
usinas em ``reports/carteiras/<nome>/`` só com o que já está gravado: o catálogo (``data/catalogo/``), os perfis achados
pelo CEG em ``usinas/*/perfil.toml`` e o ``etapa.json`` e o ``resultados.pkl`` das Análises de cada usina (com os
``indicadores_carteira``, FR-062 das Análises). Nenhuma etapa é executada e o portal não é consultado.

O ``Documento`` é a lista de blocos na ordem do contrato: capa, usinas com resultado, indicadores, ordenação, usinas sem
resultado, panorama (só com estado) e notas. O Markdown (aqui) e o PDF (``src.relatorio.carteira_pdf``) saem dele, com o
mesmo conteúdo e na mesma ordem (princípio VII). As saídas são geradas numa pasta temporária, conferidas e trocadas como
as do relatório da usina, sem ``.bak``. Nenhum texto atribui nota nem parecer (FR-049 e US10, cenário 3).
"""

from __future__ import annotations

import calendar
import pickle
import shutil
import tomllib
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import pandas as pd
from openpyxl.utils import get_column_letter

from src import pipeline
from src.analises.carteira import INDICADORES_CARTEIRA
from src.analises.conclusao import LISTAS_CONCLUSAO
from src.analises.resultados import carregar_resultados
from src.coleta.catalogo_usinas import filtrar, ler_catalogo
from src.coleta.registro import HIDRELETRICAS, RENOVAVEIS, TERMICAS, normalizar_texto
from src.comum import caminhos
from src.comum import perfil as modulo_perfil
from src.comum.formatacao import (
    SEM_VALOR,
    fmt_data,
    fmt_int,
    fmt_lista,
    fmt_mes_ano,
    fmt_num,
    fmt_pct,
    fmt_utc,
    plural,
)
from src.comum.logger import setup_logger
from src.comum.persistencia import gravar_planilha
from src.comum.regras import (
    DEFAULT_PLOT_DPI,
    LIMIAR_DESVIO_DESPACHO_MW,
    LIMIAR_GERACAO_PARADA_MW,
    MAXIMO_ITENS_CONCLUSAO,
)
from src.relatorio.carteira_figuras import NOMES_FIGURAS_CARTEIRA, gerar_figuras_carteira
from src.relatorio.carteira_pdf import gerar_pdf_carteira
from src.relatorio.etapa import PASTA_TEMPORARIA, _conferir_saidas, _trocar_saidas
from src.relatorio.fontes import COLUNAS_FONTES

logger = setup_logger("relatorio")

CODIGO_SEM_RESULTADO = 5  # catálogo ausente ou nenhuma usina do recorte com resultado (contrato)
SEM_APLICACAO = "—"  # indicador que não se aplica ao tipo, ou sem a base

MOTIVO_AGREGADO = "só agregado no ONS (ver o panorama)"
MOTIVO_SEM_PERFIL = "sem perfil"
MOTIVO_RASCUNHO = "perfil em rascunho"
MOTIVO_INVALIDO = "perfil inválido"
MOTIVO_ANALISES_AUSENTES = "Análises ausentes"
MOTIVO_ANALISES_FORMATO = "Análises em formato antigo"
MOTIVO_ANALISES_FALHA = "Análises com falha"
MOTIVO_ANALISES_DESATUALIZADAS = "Análises desatualizadas"
ORDEM_MOTIVOS: Tuple[str, ...] = (MOTIVO_AGREGADO, MOTIVO_SEM_PERFIL, MOTIVO_RASCUNHO, MOTIVO_INVALIDO,
                                  MOTIVO_ANALISES_AUSENTES, MOTIVO_ANALISES_FORMATO, MOTIVO_ANALISES_FALHA,
                                  MOTIVO_ANALISES_DESATUALIZADAS)

NOMES_COBERTURA = {"proprio": "própria", "conjunto": "conjunto", "agregado": "agregado", "ausente": "ausente"}
# Listas da conclusão na tabela da ordenação: as duas que ordenam primeiro (FR-049)
LISTAS_ORDENACAO: Tuple[str, ...] = ("possiveis_problemas", "pontos_atencao", "confirmar_agente", "verificar_campo")
TITULOS_INDICADORES: Dict[str, str] = {
    "fator_capacidade": "Fator de capacidade", "disponibilidade": "Disponibilidade", "teifa": "TEIFa",
    "energia_cortada": "Energia cortada", "atendimento_despacho": "Atendimento ao despacho",
}
# Tipos a que cada indicador se aplica (FR-062 das Análises); a coluna entra quando algum tipo do recorte é um deles
TIPOS_DO_INDICADOR: Dict[str, Tuple[str, ...]] = {
    "fator_capacidade": HIDRELETRICAS + RENOVAVEIS,
    "disponibilidade": HIDRELETRICAS + TERMICAS + RENOVAVEIS,
    "teifa": HIDRELETRICAS + TERMICAS + RENOVAVEIS,
    "energia_cortada": RENOVAVEIS,
    "atendimento_despacho": TERMICAS,
}
# Panorama: o tipo do ONS de cada tipo do catálogo (o agregado não separa UHE, PCH e CGH, nem UTE e UTN)
TIPO_ONS_DO_TIPO: Dict[str, str] = {**{t: "HIDROELETRICA" for t in HIDRELETRICAS}, **{t: "TERMICA" for t in TERMICAS},
                                    "EOL": "EOLIELETRICA", "UFV": "FOTOVOLTAICA"}
NOMES_MODALIDADES = {"PEQUENAS USINAS (MMGD)": "MMGD", "PEQUENAS USINAS (TIPO III)": "Pequenas usinas Tipo III",
                     "TIPO III": "Usinas Tipo III"}
NOMES_TIPOS_ONS = {"FOTOVOLTAICA": "fotovoltaica", "EOLIELETRICA": "eólica", "HIDROELETRICA": "hidrelétrica",
                   "TERMICA": "térmica"}
# Ordem fixa das séries do panorama (a cor de cada uma segue esta ordem, nunca a posição na figura)
ORDEM_SERIES_PANORAMA: Tuple[Tuple[str, str], ...] = (
    ("PEQUENAS USINAS (MMGD)", "FOTOVOLTAICA"), ("PEQUENAS USINAS (TIPO III)", "FOTOVOLTAICA"),
    ("PEQUENAS USINAS (TIPO III)", "EOLIELETRICA"), ("PEQUENAS USINAS (TIPO III)", "HIDROELETRICA"),
    ("PEQUENAS USINAS (TIPO III)", "TERMICA"), ("TIPO III", "HIDROELETRICA"), ("TIPO III", "TERMICA"),
)


# ---------------------------------------------------------------------------
# Recorte e situação de cada usina
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Recorte:
    """Estado, tipo ou os dois (constituição, princípio III)."""

    estado: Optional[str] = None
    tipo: Optional[str] = None

    @property
    def nome(self) -> str:
        """Nome da pasta: ``<uf>``, ``<uf>_<tipo>`` ou ``brasil_<tipo>``, em minúsculas."""
        return "_".join([self.estado or "brasil"] + ([self.tipo] if self.tipo else [])).lower()

    @property
    def rotulo(self) -> str:
        if self.estado and self.tipo:
            return f"{self.tipo} em {self.estado}"
        return self.estado or f"{self.tipo} no Brasil"


@dataclass
class UsinaComResultado:
    slug: str
    nome: str
    ceg: str
    estado: str
    tipo: str
    modalidade: str
    potencia_mw: float
    cobertura: str
    nivel_geracao: str
    indicadores: Dict[str, Any]
    analises_concluida_em: str

    @property
    def contagens(self) -> Dict[str, Any]:
        return self.indicadores.get("conclusao") or {}

    @property
    def sem_regra_avaliada(self) -> bool:
        return not self.contagens.get("regras_avaliadas")


@dataclass
class UsinaSemResultado:
    ceg: str
    nome: str
    estado: str
    tipo: str
    modalidade: str
    cobertura: str
    motivo: str
    detalhe: str = ""


@dataclass
class Carteira:
    recorte: Recorte
    catalogo_montado_em: str
    linhas_conjunto: int
    tipos: List[str]  # tipos de usina das linhas do recorte
    com_resultado: List[UsinaComResultado]  # na ordem da ordenação
    sem_resultado: List[UsinaSemResultado]  # na ordem dos motivos
    panorama: Optional[pd.DataFrame]  # linhas de agregados.csv do recorte; None na carteira só por tipo
    meses_panorama: List[str] = field(default_factory=list)
    publicacoes: Dict[str, str] = field(default_factory=dict)  # conjunto cadastral → publicação, do catalogo.json

    @property
    def total(self) -> int:
        return len(self.com_resultado) + len(self.sem_resultado)


def raiz_usinas() -> Path:
    """``usinas/`` do projeto (a mesma raiz em que ``carregar_perfil`` procura o perfil)."""
    return Path(modulo_perfil.RAIZ_PROJETO) / "usinas"


def perfis_por_ceg(raiz: Path) -> Dict[str, List[Tuple[str, str]]]:
    """{CEG normalizado: [(slug, situação)]} dos perfis em ``usinas/*/perfil.toml``; o ilegível fica de fora."""
    perfis: Dict[str, List[Tuple[str, str]]] = {}
    for arquivo in sorted(Path(raiz).glob("*/perfil.toml")) if Path(raiz).is_dir() else []:
        try:
            dados = tomllib.loads(arquivo.read_text(encoding="utf-8"))
        except (OSError, tomllib.TOMLDecodeError, UnicodeDecodeError):
            logger.warning("Perfil ilegível, fora da carteira: %s", arquivo)
            continue
        ceg = (dados.get("identificacao") or {}).get("ceg")
        if ceg:
            situacao = str((dados.get("usina") or {}).get("situacao") or "conferido")
            perfis.setdefault(normalizar_texto(ceg), []).append((arquivo.parent.name, situacao))
    return perfis


def _nivel_geracao(res: Any) -> str:
    g = (getattr(res, "niveis", None) or {}).get("geracao") or {}
    nivel = str(g.get("nivel") or "usina")
    conjunto = str(g.get("identificador_conjunto") or "")
    return f"{nivel} ({conjunto})" if conjunto else nivel


def situacao_da_usina(u: pd.Series, perfis: Dict[str, List[Tuple[str, str]]]
                      ) -> Union[UsinaComResultado, UsinaSemResultado]:
    """Com resultado (perfil conferido e válido e Análises concluídas no formato atual) ou o primeiro motivo."""
    def sem(motivo: str, detalhe: str = "") -> UsinaSemResultado:
        return UsinaSemResultado(u["ceg"], u["nom_usina"], u["id_estado"], u["tipo"], u["modalidade"],
                                 u["cobertura_resumo"], motivo, detalhe)

    if u["cobertura_resumo"] == "agregado":
        return sem(MOTIVO_AGREGADO)
    achados = perfis.get(normalizar_texto(u["ceg"]), [])
    if not achados:
        return sem(MOTIVO_SEM_PERFIL)
    if len(achados) > 1:
        return sem(MOTIVO_INVALIDO, "dois perfis com o mesmo CEG: " + fmt_lista(s for s, _ in achados))
    slug, situacao = achados[0]
    if situacao == "rascunho":
        return sem(MOTIVO_RASCUNHO, slug)
    try:
        perfil = modulo_perfil.carregar_perfil(slug)
    except modulo_perfil.PerfilInvalido:
        return sem(MOTIVO_INVALIDO, slug)
    manifesto = pipeline.ler_manifesto(slug, "analises")
    arquivo = caminhos.pasta_etapa(slug, "analises") / caminhos.ARQUIVOS_ANALISES["resultados"]
    if manifesto is None or not arquivo.is_file():
        return sem(MOTIVO_ANALISES_AUSENTES, slug)
    if manifesto.get("versao_formato") != pipeline.VERSAO_FORMATO["analises"]:
        return sem(MOTIVO_ANALISES_FORMATO, slug)
    if manifesto.get("status") == "desatualizada":
        return sem(MOTIVO_ANALISES_DESATUALIZADAS, slug)
    if manifesto.get("status") != "concluida":
        return sem(MOTIVO_ANALISES_FALHA, slug)
    try:
        res = carregar_resultados(arquivo)
    except (OSError, ValueError, EOFError, AttributeError, pickle.UnpicklingError):
        return sem(MOTIVO_ANALISES_FORMATO, slug)
    indicadores = getattr(res, "indicadores_carteira", None) or {}
    if not indicadores:
        return sem(MOTIVO_ANALISES_FORMATO, f"{slug}; sem os indicadores da carteira")
    return UsinaComResultado(
        slug=slug, nome=perfil.usina.nome, ceg=u["ceg"], estado=u["id_estado"], tipo=u["tipo"],
        modalidade=u["modalidade"] or perfil.usina.modalidade, potencia_mw=float(perfil.parametros.potencia_instalada_mw),
        cobertura=u["cobertura_resumo"], nivel_geracao=_nivel_geracao(res), indicadores=indicadores,
        analises_concluida_em=str(manifesto.get("concluida_em") or ""))


def chave_ordenacao(u: UsinaComResultado) -> Tuple[int, int, int, str]:
    """Possíveis problemas (decrescente), pontos de atenção (decrescente) e nome; sem regra avaliada, no fim."""
    c = u.contagens
    if u.sem_regra_avaliada:
        return 1, 0, 0, u.nome.casefold()
    return 0, -int(c.get("possiveis_problemas") or 0), -int(c.get("pontos_atencao") or 0), u.nome.casefold()


def _serie_panorama(modalidade: str, tipo_ons: str) -> Tuple[int, str]:
    """Posição na ordem fixa e rótulo da série (modalidade e tipo do ONS)."""
    chave = (normalizar_texto(modalidade), normalizar_texto(tipo_ons))
    posicao = ORDEM_SERIES_PANORAMA.index(chave) if chave in ORDEM_SERIES_PANORAMA else len(ORDEM_SERIES_PANORAMA)
    rotulo = f"{NOMES_MODALIDADES.get(chave[0], modalidade)} {NOMES_TIPOS_ONS.get(chave[1], tipo_ons.lower())}"
    return posicao, rotulo


def panorama_do_recorte(agregados: pd.DataFrame, recorte: Recorte) -> pd.DataFrame:
    """Linhas de ``agregados.csv`` do estado (e do tipo do ONS correspondente, com ``--tipo``), com a série."""
    a = agregados[agregados["id_estado"].str.upper() == str(recorte.estado).upper()].copy()
    if recorte.tipo:
        a = a[[normalizar_texto(t) == TIPO_ONS_DO_TIPO.get(recorte.tipo) for t in a["tipo_ons"]]]
    a["energia_mwh"] = pd.to_numeric(a["energia_mwh"], errors="coerce")
    for coluna in ("horas", "horas_ausentes"):
        a[coluna] = pd.to_numeric(a[coluna], errors="coerce").fillna(0).astype(int)
    a["horas_sem_valor"] = [_horas_do_mes(m) - h - aus for m, h, aus in zip(a["mes"], a["horas"], a["horas_ausentes"])]
    a.loc[a["horas"] == 0, "energia_mwh"] = float("nan")  # o grupo sem nenhuma hora com valor não tem energia, nem zero
    a["previsao"] = [str(v).strip().lower() == "true" for v in a["previsao"]]
    series = [_serie_panorama(m, t) for m, t in zip(a["modalidade"], a["tipo_ons"])]
    a["ordem_serie"] = [p for p, _ in series]
    a["serie"] = [rotulo + (" (previsão)" if prev else "") for (_, rotulo), prev in zip(series, a["previsao"])]
    return a.sort_values(["ordem_serie", "mes", "grupo"], kind="stable").reset_index(drop=True)


def montar_carteira(recorte: Recorte, catalogo: Dict[str, Any]) -> Carteira:
    """Recorte do catálogo, situação de cada usina, ordenação e panorama (só leitura do que está gravado)."""
    usinas = catalogo["usinas"]
    no_recorte = filtrar(usinas, recorte.estado, recorte.tipo, None)
    conjuntos = int((no_recorte["tipo"] == "conjunto").sum())
    no_recorte = no_recorte[no_recorte["tipo"] != "conjunto"]
    perfis = perfis_por_ceg(raiz_usinas())
    situacoes = [situacao_da_usina(u, perfis) for _, u in no_recorte.iterrows()]
    com = sorted((s for s in situacoes if isinstance(s, UsinaComResultado)), key=chave_ordenacao)
    sem = sorted((s for s in situacoes if isinstance(s, UsinaSemResultado)),
                 key=lambda s: (ORDEM_MOTIVOS.index(s.motivo), s.tipo, s.nome.casefold()))
    info = catalogo["json"]
    return Carteira(
        recorte=recorte, catalogo_montado_em=str(info.get("montado_em_utc") or ""), linhas_conjunto=conjuntos,
        tipos=sorted(set(no_recorte["tipo"])), com_resultado=com, sem_resultado=sem,
        panorama=panorama_do_recorte(catalogo["agregados"], recorte) if recorte.estado else None,
        meses_panorama=list((info.get("agregados") or {}).get("meses") or []),
        publicacoes={c: str(v.get("publicacao") or "") for c, v in (info.get("cadastro") or {}).items()})


# ---------------------------------------------------------------------------
# Documento: blocos na ordem do contrato
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class TabelaCarteira:
    cabecalho: List[str]
    linhas: List[List[str]]
    larguras: Tuple[float, ...]  # frações da largura útil
    numericas: Tuple[int, ...] = ()
    nota: str = ""
    fonte: str = ""  # legenda de fonte logo abaixo da tabela (princípio VII)


@dataclass(frozen=True)
class Bloco:
    tipo: str  # secao | paragrafo | tabela | figura
    texto: str = ""  # título, parágrafo ou legenda da figura
    tabela: Optional[TabelaCarteira] = None
    figura: str = ""  # chave da figura
    fonte: str = ""  # legenda de fonte logo abaixo da figura (princípio VII)


@dataclass
class Documento:
    titulo: str
    subtitulo: str
    cabecalho_pagina: str
    blocos: List[Bloco]


def texto_indicador(chave: str, indicador: Optional[Dict[str, Any]], ano_parcial: bool) -> str:
    """Valor com o detalhe que a tabela anual da usina mostra; "—" quando não se aplica ou falta a base."""
    if not indicador:
        return SEM_APLICACAO
    texto = fmt_pct(indicador["valor"])
    mes = indicador.get("mes")
    if chave == "teifa" and mes is not None and (ano_parcial or pd.Timestamp(mes).month != 12):
        texto += f" ({pd.Timestamp(mes).strftime('%m/%Y')})"
    if chave == "atendimento_despacho":
        texto += f" ({fmt_int(indicador['horas_dentro'])} de {fmt_int(indicador['horas_programadas'])} h)"
    if indicador.get("nivel", "usina") != "usina":
        texto += f" (do {indicador['nivel']})"
    return texto


def _ano(u: UsinaComResultado) -> str:
    i = u.indicadores
    return SEM_APLICACAO if i.get("ano") is None else f"{i['ano']}{'*' if i.get('ano_parcial') else ''}"


def _periodo(u: UsinaComResultado) -> str:
    p = u.indicadores.get("periodo") or {}
    if p.get("inicio") is None or p.get("fim") is None:
        return SEM_APLICACAO
    return f"{fmt_data(p['inicio'])} a {fmt_data(p['fim'])}"


PREFIXO_FONTE = "Fonte dos dados: "
PREFIXO_CALCULADO = "Calculado neste relatório a partir de: "


def origens(carteira: Carteira) -> Dict[str, Tuple[str, bool]]:
    """Origem de cada aba e da tabela que a mostra: (texto, calculada na carteira). A mesma na legenda e na FONTES."""
    catalogo = f"catálogo de usinas montado em {fmt_utc(carteira.catalogo_montado_em)}"
    analises = "resultados gravados das Análises de cada usina"
    publicacao = carteira.publicacoes.get("modalidade-usina")
    modalidade = "Modalidade das usinas (CEG de cada usina)" + (f", publicado em {fmt_data(publicacao)}"
                                                                if publicacao else "")
    meses = carteira.meses_panorama
    periodo = f" de {fmt_mes_ano(meses[0])} a {fmt_mes_ano(meses[-1])}" if meses else ""
    estado = carteira.recorte.estado
    return {
        "USINAS": (f"{modalidade}, pelo {catalogo}, com a cobertura de cada usina nos conjuntos de série do ONS; nome "
                   f"e potência: perfil de cada usina; período e nível da geração: {analises}", False),
        "INDICADORES": (f"{analises}, que vêm da série de referência e da Disponibilidade por usina (fator de capacidade "
                        "e disponibilidade), das Taxas TEIFa e TEIP (TEIFa), da Restrição com detalhe (energia cortada) e "
                        "da Geração térmica por motivo de despacho (atendimento ao despacho); o identificador e a data "
                        "de obtenção de cada conjunto estão na legenda da tabela anual do relatório de cada usina", True),
        "ORDENACAO": (f"itens da conclusão dos {analises} (aba CONCLUSAO do relatório de cada usina)", True),
        "SEM_RESULTADO": (f"cobertura de cada usina no {catalogo}; perfis em usinas/ e etapa.json das Análises de cada "
                          "usina", True),
        "PANORAMA": (f"Geração por usina (agregados de {estado} por modalidade, nível agregado), arquivos "
                     f"mensais{periodo}, lidos pelo {catalogo}" if estado else "", False),
    }


def legenda(carteira: Carteira, aba: str) -> str:
    texto, calculado = origens(carteira)[aba]
    return (PREFIXO_CALCULADO if calculado else PREFIXO_FONTE) + texto + "."


def indicadores_do_recorte(carteira: Carteira) -> List[str]:
    """Os indicadores que se aplicam a algum tipo do recorte, na ordem fixa."""
    return [c for c in INDICADORES_CARTEIRA if set(TIPOS_DO_INDICADOR[c]) & set(carteira.tipos)]


def _contagem(valor: Any) -> str:
    return SEM_APLICACAO if valor is None else fmt_int(valor)


def tabela_usinas(carteira: Carteira) -> TabelaCarteira:
    com_uf = not carteira.recorte.estado
    cabecalho = ["Usina", *(["UF"] if com_uf else []), "Tipo", "Modalidade", "Potência (MW)", "Cobertura",
                 "Nível da geração", "Período"]
    linhas = [[u.nome, *([u.estado] if com_uf else []), u.tipo, u.modalidade, fmt_num(u.potencia_mw, 1),
               NOMES_COBERTURA.get(u.cobertura, u.cobertura), u.nivel_geracao, _periodo(u)]
              for u in carteira.com_resultado]
    larguras = ((0.21, 0.05, 0.06, 0.11, 0.10, 0.10, 0.15, 0.22) if com_uf
                else (0.23, 0.06, 0.11, 0.10, 0.10, 0.17, 0.23))
    return TabelaCarteira(cabecalho, linhas, larguras, numericas=(4 if com_uf else 3,),
                          fonte=legenda(carteira, "USINAS"))


def tabela_indicadores(carteira: Carteira) -> TabelaCarteira:
    chaves = indicadores_do_recorte(carteira)
    cabecalho = ["Usina", "Ano", *[TITULOS_INDICADORES[c] for c in chaves]]
    linhas = [[u.nome, _ano(u), *[texto_indicador(c, u.indicadores.get(c), bool(u.indicadores.get("ano_parcial")))
                                  for c in chaves]] for u in carteira.com_resultado]
    resto = 0.71 / max(len(chaves), 1)
    nota = ("* Ano parcial: a série não tem ano completo. " if any(u.indicadores.get("ano_parcial")
                                                                 for u in carteira.com_resultado) else "")
    nota += f"{SEM_APLICACAO}: o indicador não se aplica ao tipo, ou a usina não tem a base."
    return TabelaCarteira(cabecalho, linhas, (0.22, 0.07, *[resto] * len(chaves)),
                          numericas=tuple(range(2, len(cabecalho))), nota=nota,
                          fonte=legenda(carteira, "INDICADORES"))


def tabela_ordenacao(carteira: Carteira) -> TabelaCarteira:
    titulos = dict(LISTAS_CONCLUSAO)
    cabecalho = ["Posição", "Usina", "Tipo", *[titulos[chave] for chave in LISTAS_ORDENACAO], "Regras avaliadas"]
    linhas = []
    for posicao, u in enumerate(carteira.com_resultado, start=1):
        c = u.contagens
        total = int(c.get("regras_avaliadas") or 0) + int(c.get("regras_nao_avaliadas") or 0)
        linhas.append([SEM_APLICACAO if u.sem_regra_avaliada else fmt_int(posicao), u.nome, u.tipo,
                       *[_contagem(c.get(chave)) for chave in LISTAS_ORDENACAO],
                       f"{fmt_int(c.get('regras_avaliadas') or 0)} de {fmt_int(total)}"])
    nota = (f"Quantidade de todos os itens de cada lista da conclusão, inclusive os que passam do limite de "
            f"{MAXIMO_ITENS_CONCLUSAO} por lista do relatório da usina.")
    return TabelaCarteira(cabecalho, linhas, (0.09, 0.23, 0.06, 0.12, 0.12, 0.12, 0.12, 0.14),
                          numericas=(0, 3, 4, 5, 6, 7), nota=nota, fonte=legenda(carteira, "ORDENACAO"))


def _motivo_texto(s: UsinaSemResultado) -> str:
    return f"{s.motivo} ({s.detalhe})" if s.detalhe else s.motivo


def tabela_sem_resultado(carteira: Carteira) -> TabelaCarteira:
    com_uf = not carteira.recorte.estado
    cabecalho = ["Usina", *(["UF"] if com_uf else []), "Tipo", "Modalidade", "Cobertura", "Motivo"]
    linhas = [[s.nome, *([s.estado] if com_uf else []), s.tipo, s.modalidade, NOMES_COBERTURA.get(s.cobertura, s.cobertura),
               _motivo_texto(s)] for s in carteira.sem_resultado]
    larguras = (0.30, 0.05, 0.06, 0.13, 0.10, 0.36) if com_uf else (0.32, 0.06, 0.13, 0.10, 0.39)
    return TabelaCarteira(cabecalho, linhas, larguras, fonte=legenda(carteira, "SEM_RESULTADO"))


def contagem_por_motivo(carteira: Carteira) -> List[Tuple[str, int]]:
    contagem: Dict[str, int] = {}
    for s in carteira.sem_resultado:
        contagem[s.motivo] = contagem.get(s.motivo, 0) + 1
    return [(m, contagem[m]) for m in ORDEM_MOTIVOS if m in contagem]


def series_do_panorama(panorama: pd.DataFrame) -> List[str]:
    """Séries presentes, na ordem fixa."""
    return list(dict.fromkeys(panorama.sort_values(["ordem_serie", "serie"], kind="stable")["serie"]))


def panorama_mensal(panorama: pd.DataFrame) -> pd.DataFrame:
    """Energia por mês e série (soma dos grupos publicados), em GWh, para a figura."""
    m = panorama.groupby(["mes", "serie"], as_index=False)["energia_mwh"].sum()
    m["gwh"] = m["energia_mwh"].fillna(0.0) / 1000.0
    m["mes"] = pd.PeriodIndex(m["mes"], freq="M").to_timestamp()
    return m


SEM_VALOR_PANORAMA = "sem valor"  # nenhum grupo da série com hora com valor no mês


def tabela_panorama(carteira: Carteira) -> TabelaCarteira:
    panorama = carteira.panorama
    series = series_do_panorama(panorama)
    energia = panorama.groupby(["mes", "serie"])["energia_mwh"].sum(min_count=1)
    meses = sorted(panorama["mes"].unique())
    linhas = [[fmt_mes_ano(mes), *[SEM_APLICACAO if (mes, s) not in energia.index
                                   else (SEM_VALOR_PANORAMA if pd.isna(energia[(mes, s)]) else fmt_num(energia[(mes, s)], 0))
                                   for s in series]] for mes in meses]
    resto = 0.86 / max(len(series), 1)
    return TabelaCarteira(["Mês", *[f"{s} (MWh)" for s in series]], linhas, (0.14, *[resto] * len(series)),
                          numericas=tuple(range(1, len(series) + 1)), fonte=legenda(carteira, "PANORAMA"))


def _horas_do_mes(mes: str) -> int:
    p = pd.Period(mes, freq="M")
    return calendar.monthrange(p.year, p.month)[1] * 24


def _horas_por_grupo(panorama: pd.DataFrame, coluna: str) -> str:
    """"N em <grupo>, de <mês> a <mês>" para cada grupo com horas na coluna; vazio se nenhum."""
    com = panorama[panorama[coluna] > 0]
    partes = []
    for grupo, g in com.groupby("grupo", sort=True):
        meses = sorted(g["mes"])
        periodo = fmt_mes_ano(meses[0]) if len(meses) == 1 else f"de {fmt_mes_ano(meses[0])} a {fmt_mes_ano(meses[-1])}"
        partes.append(f"{fmt_int(g[coluna].sum())} em {grupo} ({periodo})")
    return fmt_lista(partes)


def texto_horas_panorama(panorama: pd.DataFrame) -> str:
    esperadas = sum(_horas_do_mes(m) for m in panorama["mes"])
    return (f"Horas com valor: {fmt_int(panorama['horas'].sum())} das {fmt_int(esperadas)} horas dos grupos publicados "
            f"em cada mês. Horas sem valor: {_horas_por_grupo(panorama, 'horas_sem_valor') or 'nenhuma'}. Horas "
            f"ausentes: {_horas_por_grupo(panorama, 'horas_ausentes') or 'nenhuma'}. \"{SEM_VALOR_PANORAMA}\": nenhum "
            "grupo da série tem hora com valor no mês.")


def _titulo(recorte: Recorte) -> str:
    return f"Relatório de carteira — {recorte.rotulo}"


def montar_documento(carteira: Carteira, figuras: Dict[str, Path], data_geracao: datetime) -> Documento:
    """Os blocos do relatório, na ordem do contrato; as figuras entram só quando foram desenhadas."""
    r = carteira.recorte
    com, sem = carteira.com_resultado, carteira.sem_resultado
    blocos: List[Bloco] = []
    secao = lambda titulo: blocos.append(Bloco("secao", f"{len([b for b in blocos if b.tipo == 'secao']) + 1}. {titulo}"))  # noqa: E731
    paragrafo = lambda texto: blocos.append(Bloco("paragrafo", texto))  # noqa: E731
    tabela = lambda t: blocos.append(Bloco("tabela", tabela=t))  # noqa: E731
    n_figura = [0]

    def figura(chave: str, titulo: str, aba: str) -> None:
        if chave in figuras:
            n_figura[0] += 1
            blocos.append(Bloco("figura", f"Figura {n_figura[0]} – {titulo}", figura=chave,
                                fonte=legenda(carteira, aba)))

    escopo = f"usinas de {r.estado}" if r.estado else "usinas do Brasil"
    if r.tipo:
        escopo = f"usinas do tipo {r.tipo} " + (f"de {r.estado}" if r.estado else "do Brasil")
    paragrafo(f"Recorte: {escopo} no catálogo de usinas do ONS montado em {fmt_utc(carteira.catalogo_montado_em)}."
              + (" A linha de conjunto de usinas do catálogo não conta como usina." if carteira.linhas_conjunto == 1
                 else f" As {fmt_int(carteira.linhas_conjunto)} linhas de conjunto de usinas do catálogo não contam "
                      "como usina." if carteira.linhas_conjunto else ""))
    paragrafo(f"{fmt_int(carteira.total)} {plural(carteira.total, 'usina', 'usinas')} no recorte: {fmt_int(len(com))} "
              f"com resultado e {fmt_int(len(sem))} sem resultado. Gerado em {data_geracao.strftime('%d/%m/%Y %H:%M')}, "
              "só a partir dos resultados gravados; nenhuma etapa foi executada.")

    secao("Usinas com resultado")
    paragrafo("Usinas com perfil conferido e com as Análises concluídas, na ordem da seção 3.")
    tabela(tabela_usinas(carteira))

    secao("Indicadores comparáveis")
    paragrafo("Os indicadores da tabela anual do relatório de cada usina, no último ano completo da série, cada um "
              "do nível da usina, salvo os identificados.")
    tabela(tabela_indicadores(carteira))

    secao("Ordenação pelos itens da conclusão")
    paragrafo("Usinas ordenadas pela quantidade de itens em \"Possíveis problemas\" e, no empate, em \"Pontos de "
              "atenção\" e pelo nome.")
    tabela(tabela_ordenacao(carteira))
    sem_regra = [u.nome for u in com if u.sem_regra_avaliada]
    if sem_regra:
        paragrafo(f"A conclusão de {fmt_lista(sem_regra)} não tem base no nível da usina: nenhuma regra foi avaliada "
                  "com dado da usina, e a falta de itens não quer dizer falta de indício.")
    figura("ordenacao", "Itens da conclusão por usina, na ordem da ordenação (todos os itens de cada lista).",
           "ORDENACAO")

    secao("Usinas sem resultado")
    if sem:
        partes = [f"{motivo}, {fmt_int(n)}" for motivo, n in contagem_por_motivo(carteira)]
        paragrafo(f"{fmt_int(len(sem))} {plural(len(sem), 'usina', 'usinas')} sem resultado, por motivo: "
                  f"{'; '.join(partes)}. A cobertura é a do catálogo: própria, conjunto, agregado ou ausente (sem série "
                  "no ONS).")
        tabela(tabela_sem_resultado(carteira))
    else:
        paragrafo("Todas as usinas do recorte têm resultado.")

    if r.estado:
        secao(f"Panorama dos agregados de {r.estado}")
        p = carteira.panorama
        meses = carteira.meses_panorama
        periodo = f"de {fmt_mes_ano(meses[0])} a {fmt_mes_ano(meses[-1])}" if meses else "do catálogo"
        if p is None or p.empty:
            paragrafo(f"O ONS não publicou agregado {('do tipo ' + r.tipo + ' ') if r.tipo else ''}para {r.estado} "
                      f"{periodo}.")
        else:
            tipo_ons = f", só o tipo {NOMES_TIPOS_ONS.get(TIPO_ONS_DO_TIPO[r.tipo], '')} do ONS" if r.tipo else ""
            paragrafo(f"Geração das pequenas usinas e da MMGD de {r.estado} publicada pelo ONS de forma agregada, "
                      f"{periodo} (os 12 últimos meses completos do catálogo{tipo_ons}). Os agregados são números do "
                      "agregado, publicados por área do estado, e não são atribuídos a nenhuma usina; os valores das "
                      "usinas Tipo III são previsão do ONS.")
            tabela(tabela_panorama(carteira))
            paragrafo(texto_horas_panorama(p))
            figura("panorama", f"Energia mensal dos agregados de {r.estado} por modalidade e tipo; os valores Tipo III "
                               "são previsão do ONS.", "PANORAMA")

    blocos.append(Bloco("secao", "Notas"))
    concluidas = fmt_lista(f"{u.nome} em {fmt_utc(u.analises_concluida_em)}" for u in com)
    paragrafo(f"Fontes: o catálogo de usinas do ONS (data/catalogo/, montado em "
              f"{fmt_utc(carteira.catalogo_montado_em)}), para o recorte, a cobertura e os agregados; e os resultados "
              f"gravados das Análises de cada usina (data/usinas/<slug>/analises/), concluídas: {concluidas}.")
    paragrafo("Indicadores: fator de capacidade e disponibilidade em % da potência instalada; TEIFa publicada pelo ONS "
              "para dezembro do ano (no ano parcial, a do último mês publicado, indicado); energia cortada em % da "
              "energia estimada nas semi-horas com a marca de restrição; atendimento ao despacho: horas com programação "
              f"acima de {fmt_num(LIMIAR_GERACAO_PARADA_MW, 0)} MW em que a geração ficou a até "
              f"{fmt_num(LIMIAR_DESVIO_DESPACHO_MW, 0)} MW da programada, em % das horas com programação.")
    paragrafo("Níveis dos dados: os números de conjunto de usinas ou de agregado não são divididos entre as usinas nem "
              "atribuídos a nenhuma delas.")
    if not r.estado:
        paragrafo("Panorama dos agregados omitido: o ONS publica os agregados por área do estado, e esta carteira não "
                  "tem estado.")
    paragrafo("A carteira não atribui nota nem parecer: a ordenação só conta os itens da conclusão de cada usina, que "
              "dependem das regras avaliadas e do período de cada uma.")
    titulo = _titulo(r)
    return Documento(titulo=titulo, subtitulo="Usinas do recorte a partir dos resultados gravados das análises",
                     cabecalho_pagina=titulo, blocos=blocos)


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def _celula_md(texto: str) -> str:
    return str(texto).replace("|", "\\|").replace("*", "\\*")


def gerar_md_carteira(documento: Documento, figuras: Dict[str, Path], caminho: Path) -> Path:
    linhas: List[str] = [f"# {documento.titulo}", "", documento.subtitulo, ""]
    for bloco in documento.blocos:
        if bloco.tipo == "secao":
            linhas += [f"## {bloco.texto}", ""]
        elif bloco.tipo == "paragrafo":
            linhas += [bloco.texto, ""]
        elif bloco.tipo == "tabela":
            t = bloco.tabela
            linhas += ["| " + " | ".join(_celula_md(c) for c in t.cabecalho) + " |",
                       "| " + " | ".join("---:" if i in t.numericas else "---" for i in range(len(t.cabecalho))) + " |"]
            linhas += ["| " + " | ".join(_celula_md(c) for c in linha) + " |" for linha in t.linhas] + [""]
            linhas += [_celula_md(t.fonte), ""]  # a legenda de fonte logo abaixo da tabela (princípio VII)
            if t.nota:
                linhas += [_celula_md(t.nota), ""]
        elif bloco.tipo == "figura":
            nome = Path(figuras[bloco.figura]).name
            linhas += [f"![{bloco.texto}]({caminhos.ARQUIVOS_CARTEIRA['figuras']}/{nome})", "", f"*{bloco.texto}*", "",
                       _celula_md(bloco.fonte), ""]
    destino = Path(caminho)
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text("\n".join(linhas), encoding="utf-8")
    return destino


# ---------------------------------------------------------------------------
# Planilha
# ---------------------------------------------------------------------------


def _detalhe_indicador(chave: str, indicador: Optional[Dict[str, Any]]) -> str:
    if not indicador:
        return ""
    if chave == "teifa" and indicador.get("mes") is not None:
        return f"mês {pd.Timestamp(indicador['mes']).strftime('%m/%Y')}"
    if chave == "atendimento_despacho":
        return f"{indicador['horas_dentro']} de {indicador['horas_programadas']} horas com programação"
    return ""


def abas_planilha(carteira: Carteira) -> Dict[str, pd.DataFrame]:
    """Abas USINAS, INDICADORES, ORDENACAO, SEM_RESULTADO, PANORAMA e FONTES (data-model, seção 6.5)."""
    listas = [chave for chave, _ in LISTAS_CONCLUSAO]
    usinas, indicadores, ordenacao = [], [], []
    for posicao, u in enumerate(carteira.com_resultado, start=1):
        i, c = u.indicadores, u.contagens
        periodo = i.get("periodo") or {}
        linha: Dict[str, Any] = {
            "slug": u.slug, "nome": u.nome, "ceg": u.ceg, "estado": u.estado, "tipo": u.tipo,
            "modalidade": u.modalidade, "potencia_mw": u.potencia_mw, "cobertura": u.cobertura,
            "nivel_geracao": u.nivel_geracao, "periodo_inicio": periodo.get("inicio"), "periodo_fim": periodo.get("fim"),
            "ano": i.get("ano"), "ano_parcial": bool(i.get("ano_parcial")),
        }
        for chave in INDICADORES_CARTEIRA:
            ind = i.get(chave) or {}
            linha[f"{chave}_pct"] = ind.get("valor")
            linha[f"{chave}_nivel"] = ind.get("nivel", "")
            indicadores.append({"slug": u.slug, "nome": u.nome, "tipo": u.tipo, "ano": i.get("ano"),
                                "ano_parcial": bool(i.get("ano_parcial")), "indicador": chave,
                                "valor": ind.get("valor"), "nivel": ind.get("nivel", ""),
                                "unidade": ind.get("unidade", ""), "detalhe": _detalhe_indicador(chave, ind)})
        contagens = {chave: c.get(chave) for chave in listas}
        regras = {"regras_avaliadas": c.get("regras_avaliadas"), "regras_nao_avaliadas": c.get("regras_nao_avaliadas")}
        usinas.append({**linha, **contagens, **regras, "analises_concluida_em": u.analises_concluida_em})
        ordenacao.append({"posicao": None if u.sem_regra_avaliada else posicao, "slug": u.slug, "nome": u.nome,
                          "tipo": u.tipo, **contagens, **regras})
    inteiros = [*listas, "regras_avaliadas", "regras_nao_avaliadas", "ano"]
    tabelas = {"USINAS": pd.DataFrame(usinas), "INDICADORES": pd.DataFrame(indicadores),
               "ORDENACAO": pd.DataFrame(ordenacao)}
    for nome in ("USINAS", "ORDENACAO", "INDICADORES"):
        t = tabelas[nome]
        for coluna in [c for c in [*inteiros, "posicao"] if c in t.columns]:
            t[coluna] = pd.array([None if v is None or pd.isna(v) else int(v) for v in t[coluna]], dtype="Int64")
    sem = pd.DataFrame([{"ceg": s.ceg, "nome": s.nome, "estado": s.estado, "tipo": s.tipo, "modalidade": s.modalidade,
                         "cobertura": s.cobertura, "motivo": s.motivo, "detalhe": s.detalhe}
                        for s in carteira.sem_resultado],
                       columns=["ceg", "nome", "estado", "tipo", "modalidade", "cobertura", "motivo", "detalhe"])
    if carteira.panorama is not None:
        panorama = carteira.panorama.drop(columns=["ordem_serie"])
    else:
        panorama = pd.DataFrame({"motivo": ["Panorama omitido: o ONS publica os agregados por área do estado, e a "
                                            "carteira só por tipo não tem estado."]})
    # FONTES com as colunas da aba da usina (D22): a origem é a mesma da legenda de cada tabela
    fontes = pd.DataFrame([{"aba": aba, "conjuntos_origem": texto or SEM_VALOR, "conferencias": SEM_VALOR,
                            "sem_outra_fonte": SEM_VALOR, "calculado_no_relatorio": "sim" if calculado else "não"}
                           for aba, (texto, calculado) in origens(carteira).items()], columns=COLUNAS_FONTES)
    return {**tabelas, "SEM_RESULTADO": sem, "PANORAMA": panorama, "FONTES": fontes}


def _formatar_planilha(writer: pd.ExcelWriter) -> None:
    for folha in writer.book.worksheets:
        folha.freeze_panes = "A2"
        for indice, coluna in enumerate(folha.iter_cols(min_row=1, max_row=min(folha.max_row, 200)), start=1):
            largura = max((len(str(c.value)) for c in coluna if c.value is not None), default=8)
            folha.column_dimensions[get_column_letter(indice)].width = min(max(largura + 2, 8), 60)


# ---------------------------------------------------------------------------
# Gravação e comando
# ---------------------------------------------------------------------------


def gravar_carteira(carteira: Carteira, pasta: Path, data_geracao: Optional[datetime] = None) -> List[Path]:
    """Gera as saídas numa pasta temporária, confere e troca as da execução anterior (sem ``.bak``)."""
    arq = caminhos.ARQUIVOS_CARTEIRA
    temporaria = Path(pasta) / PASTA_TEMPORARIA
    if temporaria.exists():
        shutil.rmtree(temporaria)
    temporaria.mkdir(parents=True)
    gerado_em = data_geracao or datetime.now()
    try:
        ordenacao = [(u.nome, u.contagens) for u in carteira.com_resultado]
        p = carteira.panorama
        mensal = panorama_mensal(p) if p is not None and len(p) else None
        series = series_do_panorama(p) if p is not None and len(p) else []
        figuras = gerar_figuras_carteira(ordenacao, mensal, series,
                                         f"Agregados de {carteira.recorte.estado}: energia mensal por modalidade e tipo",
                                         temporaria / arq["figuras"], DEFAULT_PLOT_DPI)
        documento = montar_documento(carteira, figuras, gerado_em)
        md = gerar_md_carteira(documento, figuras, temporaria / arq["md"])
        pdf = gerar_pdf_carteira(documento, figuras, temporaria / arq["pdf"], invariante=data_geracao is not None)
        xlsx = temporaria / arq["xlsx"]
        gravar_planilha(abas_planilha(carteira), xlsx, _formatar_planilha)
        novos = [pdf, md, xlsx, *sorted(figuras.values())]
        _conferir_saidas(novos)
        arquivos = _trocar_saidas(temporaria, Path(pasta), novos)
    finally:
        shutil.rmtree(temporaria, ignore_errors=True)
    for nome in NOMES_FIGURAS_CARTEIRA.values():  # figura da execução anterior que agora não tem base
        antiga = Path(pasta) / arq["figuras"] / nome
        if antiga not in arquivos and antiga.exists():
            antiga.unlink()
            logger.info("Figura de execução anterior removida: %s", nome)
    return arquivos


def _sugestao_slug(carteira: Carteira) -> str:
    for s in carteira.sem_resultado:
        if s.motivo.startswith("Análises") and s.detalhe:
            return s.detalhe.split(";")[0]
    return "<slug>"


def executar_carteira(estado: Optional[str] = None, tipo: Optional[str] = None,
                      data_geracao: Optional[datetime] = None) -> int:
    """Executa o comando ``carteira`` e devolve o código de saída do contrato (0, 1, 2 ou 5)."""
    if not estado and not tipo:
        logger.error("carteira: use --estado, --tipo ou os dois.")
        return pipeline.CODIGO_DADO_NAO_OBTIDO
    recorte = Recorte(estado.upper() if estado else None, tipo)
    catalogo = ler_catalogo(caminhos.pasta_catalogo())
    if catalogo is None:
        logger.error("O catálogo de usinas não foi montado. Execute antes: python -m src usinas")
        return CODIGO_SEM_RESULTADO
    try:
        carteira = montar_carteira(recorte, catalogo)
        if not carteira.com_resultado:
            logger.error("Nenhuma usina do recorte %s tem as Análises concluídas. Execute antes: python -m src "
                         "completo --usina %s", recorte.rotulo, _sugestao_slug(carteira))
            return CODIGO_SEM_RESULTADO
        pasta = caminhos.pasta_carteira(recorte.nome)
        gravar_carteira(carteira, pasta, data_geracao)
    except KeyboardInterrupt:
        logger.error("Carteira %s interrompida pelo usuário.", recorte.nome)
        return pipeline.CODIGO_ERRO
    except Exception as exc:
        logger.exception("Carteira %s: erro inesperado: %s", recorte.nome, exc)
        return pipeline.CODIGO_ERRO
    logger.info("Carteira %s: %d usinas com resultado e %d sem resultado. Pasta: reports/%s/%s/", recorte.nome,
                len(carteira.com_resultado), len(carteira.sem_resultado), caminhos.PASTA_CARTEIRAS, recorte.nome)
    return pipeline.CODIGO_SUCESSO
